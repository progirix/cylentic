#!/usr/bin/env python3
"""Intègre chapitres 4 et 5 (corrigés) dans le rapport final, sans régénération totale."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

FINAL = "/workspace/docs/Rapport_Cylentic_Carcasse_Chapitre3.docx"


def set_run_font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def insert_paragraph_after(paragraph, text, **kwargs):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    run = new_para.add_run(text)
    set_run_font(
        run,
        size=kwargs.get("size", 12),
        bold=kwargs.get("bold", False),
        italic=kwargs.get("italic", False),
    )
    align = kwargs.get("align", "justify")
    if align == "center":
        new_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "left":
        new_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        new_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = new_para.paragraph_format
    pf.space_after = Pt(kwargs.get("space_after", 8))
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if kwargs.get("first_line", True) and align == "justify":
        pf.first_line_indent = Cm(0.75)
    return new_para


def insert_heading_after(paragraph, text, level=2):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    # use heading style if available
    try:
        new_para.style = f"Heading {min(level, 3)}"
    except Exception:
        pass
    run = new_para.add_run(text)
    set_run_font(run, size=14 if level <= 2 else 12, bold=True)
    new_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = new_para.paragraph_format
    pf.space_before = Pt(12)
    pf.space_after = Pt(8)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    return new_para


def find_para(doc, pred, prefer_heading=False):
    matches = [p for p in doc.paragraphs if pred(p.text.strip())]
    if not matches:
        return None
    if prefer_heading:
        headed = [p for p in matches if p.style and str(p.style.name).startswith("Heading")]
        if headed:
            return headed[-1]
    return matches[-1]


def clear_until(doc, heading_p, stops):
    to_remove = []
    started = False
    for p in list(doc.paragraphs):
        if p._p is heading_p._p:
            started = True
            continue
        if not started:
            continue
        t = p.text.strip()
        if any(t.startswith(sp) for sp in stops):
            break
        to_remove.append(p)
    for p in to_remove:
        parent = p._element.getparent()
        if parent is not None:
            parent.remove(p._element)


def replace_paragraph_text(p, text):
    # clear runs
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    set_run_font(run, size=12)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Cm(0.75)
    pf.space_after = Pt(8)


def fill_with_blocks(doc, chapter_heading, stops, blocks):
    """
    blocks: list of ('h2'|'h3'|'p'|'fig', text, opts?)
    Removes everything under chapter until stops, rebuilds headings + content.
    """
    heading_p = find_para(
        doc,
        lambda t, h=chapter_heading: t.startswith(h),
        prefer_heading=True,
    )
    if heading_p is None:
        raise SystemExit(f"Missing {chapter_heading}")
    clear_until(doc, heading_p, stops)
    anchor = heading_p
    for kind, text, *rest in blocks:
        opts = rest[0] if rest else {}
        if kind == "h2":
            anchor = insert_heading_after(anchor, text, level=2)
        elif kind == "h3":
            anchor = insert_heading_after(anchor, text, level=3)
        elif kind == "fig":
            number, title, desc = text
            anchor = insert_paragraph_after(
                anchor,
                f"[Emplacement de la figure {number}]",
                italic=True,
                align="center",
                first_line=False,
                space_after=4,
            )
            anchor = insert_paragraph_after(
                anchor,
                desc,
                italic=True,
                align="center",
                size=11,
                first_line=False,
                space_after=4,
            )
            anchor = insert_paragraph_after(
                anchor,
                f"Figure {number} : {title}",
                size=11,
                align="center",
                first_line=False,
                space_after=12,
            )
        else:
            anchor = insert_paragraph_after(anchor, text, **opts)
    return heading_p


def main():
    doc = Document(FINAL)

    # ---- Fix périmètre MVP (exports maintenant livrés) ----
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("Sont exclus de ce périmètre"):
            replace_paragraph_text(
                p,
                "Sont exclus de ce périmètre : le support d'autres langages que Python, "
                "la facturation réelle des établissements, la consultation en direct du "
                "code pendant l'épreuve, et l'activation de compte par email. Le MVP "
                "s'appuie sur des comptes créés par l'administration, un changement de "
                "mot de passe imposé si nécessaire, et des exports de résultats déjà "
                "disponibles pour le professeur (Excel, PDF) ainsi qu'un export de "
                "présence (CSV).",
            )
        if t.startswith("Le projet assume un périmètre volontairement limité"):
            replace_paragraph_text(
                p,
                "Le projet assume un périmètre volontairement limité pour sa première "
                "version. Le MVP inclut : le langage Python, l'authentification avec "
                "rôles (Super Admin, administrateur d'établissement, professeur, "
                "étudiant), la création et la publication d'examens, la salle "
                "d'attente, l'environnement de composition sécurisé, les mécanismes de "
                "contrôle navigateur, la correction automatique par tests unitaires, "
                "un module QCM de base, la gestion de plusieurs établissements, "
                "l'onboarding administrateur, la promotion de rentrée, le contact "
                "depuis la landing page, ainsi que les exports de résultats et de "
                "présence.",
            )

    # Soften Redis/workers as architecture cible vs livré
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("MySQL conserve la mémoire durable"):
            replace_paragraph_text(
                p,
                "MySQL conserve la mémoire durable : sujets, copies, scores, journaux. "
                "Redis et des workers de file sont prévus dans l'architecture cible "
                "pour absorber des pics de correction et enrichir le temps réel. Dans "
                "le MVP livré, ces briques ne sont pas encore industrialisées : le "
                "suivi live repose sur du polling HTTP, et la correction automatique "
                "est déclenchée dans le parcours de soumission.",
            )
        if "héberge Next.js et les workers" in t and t.startswith("Le diagramme de déploiement"):
            replace_paragraph_text(
                p,
                "Le diagramme de déploiement traduit cette architecture en nœuds "
                "d'exécution. Le poste utilisateur héberge le navigateur. Le serveur "
                "d'application héberge Next.js. MySQL porte la persistance. Redis et "
                "les workers apparaissent comme perspectives d'industrialisation. "
                "L'hôte d'exécution isolée accueille le conteneur sandbox. En "
                "production, le lien navigateur application est prévu en HTTPS. En "
                "démonstration salle, le même produit peut tourner en HTTP sur le "
                "réseau local.",
            )

    ch4_blocks = [
        (
            "p",
            "Une fois le cadrage et la conception posés, la réalisation a pour objet "
            "de transformer le cahier des charges en un produit exécutable. L'équipe "
            "a cherché un dépôt unique, déployable localement, capable de faire vivre "
            "les rôles de la plateforme et surtout de faire passer un examen du début "
            "à la fin : création, publication, composition sous contraintes, "
            "correction, puis exports.",
        ),
        ("h2", "4.1 Environnement de travail"),
        (
            "p",
            "Le projet repose sur Node.js et npm, avec Next.js 16 en App Router pour "
            "les pages et les API. Prisma 7 s'appuie sur MySQL pour le modèle de "
            "données. Docker sert à isoler l'exécution du code Python des étudiants. "
            "L'édition et le suivi se font principalement dans VS Code ou Cursor. Les "
            "essais de passation et de contrôle navigateur ont été menés en priorité "
            "sur Chrome et Edge, navigateurs compatibles avec l'API Fullscreen.",
        ),
        (
            "p",
            "Pour installer et démarrer le projet en local, la séquence habituelle est "
            "npm install, puis npm run db:setup, puis npm run dev. En démonstration "
            "multi postes sur le réseau local, un build de production est préférable "
            "(npm run build, puis npx next start -H 0.0.0.0 -p 3000). Les variables "
            "critiques restent DATABASE_URL, JWT_SECRET et JWT_EXPIRES_IN. Le drapeau "
            "FORCE_SECURE_COOKIES n'est utile que derrière un reverse proxy HTTPS.",
        ),
        ("h2", "4.2 Choix technologiques"),
        (
            "p",
            "Le choix d'un monolithe Next.js centralise pages et routes API dans un "
            "seul déploiement, ce qui simplifie le jour J. Prisma apporte un typage "
            "fort et un schéma déclaratif, au prix de quelques subtilités de "
            "sérialisation, notamment pour les décimaux côté composants clients. "
            "L'authentification par JWT en cookie httpOnly reste adaptée à un MVP et "
            "évite une session serveur lourde. Monaco Editor offre une expérience "
            "proche d'un IDE moderne.",
        ),
        (
            "p",
            "L'exécution du code étudiant dans un conteneur Docker vise une isolation "
            "réelle. Zod valide les entrées API. Le temps réel a volontairement été "
            "traité par polling HTTP plutôt que par WebSocket et Redis, afin de livrer "
            "d'abord un comportement fiable sans infrastructure supplémentaire. Redis "
            "et les workers restent donc des perspectives d'industrialisation, et non "
            "des dépendances du parcours critique actuel.",
        ),
        ("h2", "4.3 Organisation du code et architecture applicative"),
        (
            "p",
            "La règle générale a été de limiter la logique métier dans les pages. Les "
            "pages authentifient, orchestrent et affichent. Les services encapsulent "
            "les invariants : limites d'administrateurs, verrous d'examen, correction, "
            "promotion de rentrée, exports, etc.",
        ),
        (
            "p",
            "Côté présentation, les routes publiques regroupent landing, connexion et "
            "inscription d'établissement. Les espaces authentifiés sont séparés par "
            "rôle. Les API suivent une convention JSON homogène, avec un booléen "
            "success et, selon le cas, un objet data ou un message d'erreur. Les "
            "services mobilisés couvrent notamment l'authentification, les "
            "utilisateurs, les examens, la participation, la correction, la sandbox, "
            "le billing des plans, la promotion, le contact landing, les exports et "
            "les tableaux de bord.",
        ),
        ("h2", "4.4 Réalisation des modules principaux"),
        ("h3", "4.4.1 Authentification et sessions"),
        (
            "p",
            "À la connexion, le rôle n'est pas choisi manuellement : il est déduit du "
            "préfixe d'identifiant (SADM-, ADM-, PROF-, ETU-). Le JWT signe les "
            "informations utiles (identifiant technique, identité affichable, rôle, "
            "établissement, obligation de changer le mot de passe) et est stocké dans "
            "le cookie de session, en httpOnly, avec SameSite=Lax.",
        ),
        (
            "p",
            "Un point sensible a été corrigé pendant la réalisation. Avec next start, "
            "l'environnement passe en production et le cookie était initialement "
            "marqué Secure. Or, en accès salle via une adresse HTTP locale, le "
            "navigateur refuse ce cookie. La connexion semblait alors échouer pour "
            "tous les profils. La solution livrée active Secure uniquement lorsque la "
            "requête est réellement en HTTPS, ou si FORCE_SECURE_COOKIES est "
            "explicitement forcé.",
        ),
        ("h3", "4.4.2 Module administrateur d'établissement"),
        (
            "p",
            "Après inscription, l'administrateur est guidé dans un onboarding : année "
            "académique, classes, premiers utilisateurs. Tant que ces étapes ne sont "
            "pas complètes, le tableau de bord peut rediriger vers le parcours "
            "d'installation. Les tableaux de bord présentent ensuite population, "
            "examens et activité. Les étudiants peuvent être créés manuellement ou "
            "importés par CSV, avec un rapport d'erreurs ligne à ligne. Les "
            "professeurs reçoivent un identifiant PROF-* et un mot de passe par "
            "défaut. Les administrateurs sont plafonnés à deux comptes actifs : on ne "
            "peut pas s'auto désactiver, et le système refuse de laisser "
            "l'établissement sans administrateur actif.",
        ),
        (
            "p",
            "La promotion de rentrée a aussi été livrée. Le parcours propose le choix "
            "ou la création de l'année cible, l'association des classes sources vers "
            "les classes de destination, puis la confirmation. La transaction met à "
            "jour les profils étudiants, active la nouvelle année et peut archiver "
            "l'année précédente.",
        ),
        ("h3", "4.4.3 Module professeur"),
        (
            "p",
            "Le cycle de vie d'un examen reste celui conçu au chapitre 3. Le "
            "professeur crée un brouillon, compose les exercices (code Python avec "
            "tests et/ou QCM), puis publie. La publication génère un code d'accès "
            "affichable aussi en mode présentation pour le vidéoprojecteur. Pendant "
            "l'épreuve, le suivi live montre les statuts. Après l'examen, le "
            "professeur consulte les copies, peut ajuster une note manuelle et "
            "exporter les résultats (Excel, PDF) ainsi que la présence (CSV).",
        ),
        (
            "p",
            "Plusieurs capacités ont étendu le MVP en cours de projet. Le planning "
            "peut être modifié sur un examen publié, avec synchronisation côté "
            "composition. Le contenu reste modifiable tant que la date de début n'est "
            "pas atteinte, via une règle partagée de verrouillage. La duplication "
            "crée un nouveau brouillon en reprenant la structure sans recopier les "
            "participations.",
        ),
        ("h3", "4.4.4 Module étudiant et parcours de composition"),
        (
            "p",
            "Le parcours étudiant est volontairement explicite. La première page "
            "présente les consignes et exige le plein écran. Tant que l'API "
            "Fullscreen est refusée ou indisponible, l'étudiant ne peut pas avancer. "
            "Une fois le plein écran accepté, il entre en salle d'attente. Un compte "
            "à rebours est synchronisé avec le serveur via un polling sur l'API de "
            "session. L'énoncé reste masqué. Dès cette étape, l'heure de connexion "
            "et l'adresse IP sont enregistrées.",
        ),
        (
            "fig",
            (
                "4.1",
                "Salle d'attente avant le démarrage de l'examen",
                "Insérer la capture de la salle d'attente (compte à rebours, énoncé masqué).",
            ),
        ),
        (
            "p",
            "Au démarrage, la page de composition affiche le menu des exercices, "
            "l'énoncé, l'éditeur Monaco pour Python, le bouton d'exécution, le timer "
            "global et la soumission. Un autosave périodique, complété par le "
            "stockage local du navigateur, limite les pertes en cas de coupure "
            "brève. Un second polling permet de mettre à jour la fin d'examen si le "
            "professeur modifie le planning en cours de session.",
        ),
        (
            "fig",
            (
                "4.2",
                "Environnement de composition pendant l'épreuve",
                "Insérer la capture de la page de composition (IDE, timer, exercices).",
            ),
        ),
        (
            "p",
            "Les contrôles de passation signalent au serveur les sorties de plein "
            "écran, les changements d'onglet et certains collages. Au second incident "
            "de sortie du plein écran, la participation est exclue et redirigée vers "
            "une confirmation avec motif. La soumission peut être manuelle, "
            "automatique à l'expiration du timer, ou forcée par exclusion. Une fois "
            "la participation close, aucune reconnexion utile n'est possible. Aucun "
            "score n'est montré à l'étudiant.",
        ),
        ("h3", "4.4.5 Sandbox et correction automatique"),
        (
            "p",
            "Le service sandbox exécute le code Python dans un conteneur éphémère, "
            "avec timeout et sans accès réseau. En développement, un repli local peut "
            "exister si Docker est absent ; ce mode n'est pas recommandé pour une "
            "salle réelle. La correction automatique parcourt chaque test unitaire, "
            "exécute le programme, normalise la sortie, puis compare à la sortie "
            "attendue. Le poids de chaque test intervient dans le score automatique. "
            "Le professeur peut ensuite imposer un score manuel, qui prend le dessus "
            "lorsqu'il est renseigné.",
        ),
        ("h3", "4.4.6 Landing page et contact"),
        (
            "p",
            "La landing présente le produit. Le formulaire de contact appelle une API "
            "dédiée, enregistre une notification de type landing.contact et rend "
            "cette demande consultable côté Super Admin.",
        ),
        ("h2", "4.5 Points techniques difficiles"),
        (
            "p",
            "Plusieurs problèmes concrets ont marqué la réalisation. Le cookie Secure "
            "en HTTP bloquait les connexions sur le réseau local ; il a fallu le "
            "conditionner à la présence réelle de HTTPS. Les objets Decimal de Prisma "
            "provoquaient un crash lorsqu'ils étaient passés directement à des "
            "composants clients ; la réponse a été de sérialiser les points en "
            "nombres côté serveur. Le timer étudiant ne suivait pas les changements "
            "de planning tant que la session n'était chargée qu'une seule fois ; un "
            "polling de synchronisation a réglé ce point. Le verrouillage trop strict "
            "des exercices empêchait toute modification d'un examen publié avant le "
            "démarrage ; la règle canModifyExamExercises a clarifié le comportement. "
            "Enfin, le passage d'icônes depuis des fichiers serveur vers des "
            "composants clients a exigé un mapping par chaînes côté client.",
        ),
        ("h2", "4.6 Déploiement"),
        (
            "p",
            "Pour une démonstration en salle, le scénario recommandé est simple : "
            "produire un build, démarrer avec next start en écoutant sur toutes les "
            "interfaces, puis accéder via http://adresse-de-l-hote:3000 depuis les "
            "postes étudiants. Il faut parfois ouvrir le port 3000 dans le pare-feu "
            "de la machine hôte. En production derrière HTTPS, le cookie Secure peut "
            "être activé normalement, avec ou sans FORCE_SECURE_COOKIES selon le "
            "proxy.",
        ),
        ("h2", "4.7 Conclusion partielle"),
        (
            "p",
            "Le MVP livre aujourd'hui une chaîne complète, suffisamment intégrée pour "
            "la démonstration et pour des usages pilotes. Les priorités d'ingénierie "
            "ont privilégié la fiabilité des parcours critiques (authentification, "
            "examen, contrôles de passation, correction, exports) plutôt qu'une "
            "infrastructure temps réel complète. Redis, les workers et le multi "
            "langage restent volontairement dans les perspectives.",
        ),
    ]

    ch5_blocks = [
        (
            "p",
            "Ce chapitre présente la campagne de validation du MVP, les résultats "
            "observés, les anomalies corrigées en cours de route, les limites "
            "assumées et les perspectives de poursuite.",
        ),
        ("h2", "5.1 Stratégie de tests"),
        (
            "p",
            "La campagne cherchait à confirmer les exigences fonctionnelles critiques, "
            "à contrôler les règles métier, à vérifier la non régression technique via "
            "le build TypeScript, et à s'assurer du bon fonctionnement multi postes en "
            "HTTP local. Les tests unitaires automatisés exhaustifs n'ont pas été "
            "l'axe principal sous contrainte de temps. L'équipe a privilégié des "
            "scénarios bout en bout, plus représentatifs d'un livrable démontrable.",
        ),
        (
            "p",
            "La validation a donc combiné le build Next.js, des parcours manuels rôle "
            "par rôle, des essais de bornes sur les règles métier, des tests multi "
            "appareils sur le réseau local, et des vérifications après chaque "
            "correctif majeur (cookie Secure, Decimal, synchronisation du timer). Les "
            "jeux de données s'appuyaient sur un établissement de démonstration, des "
            "comptes types, des examens dans plusieurs états et des CSV volontairement "
            "valides ou invalides.",
        ),
        ("h2", "5.2 Jeux de tests et résultats"),
        ("h3", "Authentification et accès"),
        (
            "p",
            "Les essais de connexion montrent un comportement stable pour les "
            "différents rôles lorsque les identifiants sont corrects. Un mauvais mot "
            "de passe est rejeté. Un étudiant sans code d'examen ne peut pas démarrer "
            "le parcours d'épreuve. Le middleware empêche un rôle d'accéder aux "
            "espaces d'un autre. La déconnexion vide correctement le cookie. L'accès "
            "depuis un autre appareil en HTTP fonctionne après le correctif Secure.",
        ),
        ("h3", "Administration"),
        (
            "p",
            "La création de classes et d'étudiants manuels fonctionne, avec génération "
            "d'identifiants. L'import CSV crée les lignes valides et rejette clairement "
            "les classes inconnues ou les emails incorrects, en laissant un rapport "
            "exploitable. La limite de deux administrateurs actifs est respectée, "
            "tout comme l'interdiction de se désactiver soi-même. La promotion de "
            "rentrée met à jour les profils et l'année active. Le formulaire de "
            "contact de la landing enregistre une notification visible côté Super "
            "Admin.",
        ),
        ("h3", "Examens côté professeur"),
        (
            "p",
            "Un examen sans exercice ne peut pas être publié. Avec du contenu, la "
            "publication génère le code d'accès et permet le mode présentation. La "
            "duplication produit un nouveau brouillon complet. L'édition du planning "
            "sur un examen publié se répercute en salle d'attente et en composition. "
            "L'ajout d'exercices reste possible avant le démarrage et est refusé "
            "après. Le suivi live, les exports et l'ajustement manuel des notes se "
            "comportent conformément aux règles prévues.",
        ),
        ("h3", "Parcours étudiant et contrôles de passation"),
        (
            "p",
            "Un code d'examen invalide est refusé, avec protection contre les "
            "tentatives excessives. Sans plein écran, l'étudiant reste bloqué aux "
            "consignes. Au démarrage, la bascule vers la composition est automatique. "
            "L'exécution sandbox, l'autosave, la navigation entre exercices et la "
            "soumission manuelle ou automatique au timer fonctionnent. Les incidents "
            "de plein écran incrémentent un compteur et conduisent à l'exclusion au "
            "second incident. Le changement d'onglet est journalisé. Après "
            "soumission, la reconnexion est bloquée.",
        ),
        (
            "p",
            "Au total, la campagne manuelle ciblée a porté sur trente neuf cas "
            "critiques (authentification, administration, professeur, étudiant). Tous "
            "sont passés après les correctifs majeurs. Ce résultat doit être lu comme "
            "une validation fonctionnelle du périmètre MVP, et non comme une couverture "
            "automatisée exhaustive.",
        ),
        ("h2", "5.3 Discussion des résultats"),
        (
            "p",
            "Le MVP permet aujourd'hui un parcours d'examen crédible, de la salle "
            "d'attente jusqu'à la soumission sous contraintes. La correction "
            "automatique des tests unitaires réduit le temps de traitement pour le "
            "professeur sans lui retirer le contrôle pédagogique. Les incidents "
            "horodatés, la présence numérique et le journal d'activité offrent une "
            "traçabilité utile. L'outil est exploitable en salle sur réseau local. "
            "L'administration scolaire et la présence produit (landing, contact, "
            "espaces par rôle) complètent un livrable démontrable de bout en bout.",
        ),
        (
            "p",
            "Trois anomalies ont particulièrement pesé avant d'être corrigées. Le "
            "cookie Secure bloquait toutes les connexions locales. Les props Decimal "
            "cassaient des composants clients. Le timer de composition n'évoluait pas "
            "après une modification professeur sans polling de synchronisation. Ces "
            "points confirment que la valeur du MVP tient autant aux parcours qu'à la "
            "robustesse des détails d'exécution en salle.",
        ),
        ("h2", "5.4 Limites et risques résiduels"),
        (
            "p",
            "La sécurité navigateur reste une dissuasion et une détection, non une "
            "isolation absolue de poste. Un second appareil, une capture d'écran "
            "externe ou une machine virtuelle hors contrôle ne sont pas bloqués "
            "techniquement. Le langage supporté est encore limité à Python. La "
            "sandbox dépend de Docker. Les workers et la file de correction ne sont "
            "pas industrialisés. Le temps réel repose sur du polling, acceptable pour "
            "des promotions modestes mais à revoir sous forte charge. L'activation "
            "par email, le paiement réel, une intégration Moodle profonde, certains "
            "raffinements QCM et une suite de tests bout en bout automatisés restent "
            "hors du livrable actuel ou partiels.",
        ),
        ("h2", "5.5 Perspectives"),
        (
            "p",
            "La feuille de route naturelle commence par l'élargissement des langages "
            "(Java, C, C++), l'activation des comptes par email et l'introduction de "
            "Redis avec un canal temps réel plus robuste. Les QCM avancés, la "
            "facturation réelle, une file de correction capable d'absorber des pics "
            "de soumissions, puis une API publique ou Moodle, constituent les étapes "
            "suivantes. Une interface dédiée au registre de salle pour les "
            "surveillants fermerait un chaînon encore humain. Toute surveillance "
            "vidéo, en revanche, devra d'abord trancher des questions éthiques et de "
            "protection des données avant d'être envisagée.",
        ),
        (
            "p",
            "Quatre recommandations pragmatiques ressortent pour la suite. D'abord, "
            "prioriser une suite bout en bout automatisée sur le trajet connexion, "
            "publication, composition et soumission. Ensuite, introduire Redis dès "
            "que le live doit dépasser quelques dizaines d'étudiants simultanés. "
            "Préparer ensuite une procédure salle simple (adresse, pare-feu, "
            "navigateurs supportés). Enfin, cadrer tôt le positionnement juridique "
            "avant toute dérive vers une surveillance intrusive.",
        ),
        ("h2", "5.6 Conclusion partielle"),
        (
            "p",
            "La validation confirme que le MVP Cylentic est fonctionnellement "
            "démontrable et techniquement cohérent sur son périmètre. Les limites "
            "restantes relèvent surtout d'arbitrages de priorité (infrastructure "
            "temps réel, multi langages, automatisation des tests), et non d'un "
            "défaut d'architecture fondamental.",
        ),
    ]

    # Convert tuples of length 2 to triples for fill
    def normalize(blocks):
        out = []
        for b in blocks:
            if len(b) == 2:
                out.append((b[0], b[1], {}))
            else:
                out.append(b)
        return out

    fill_with_blocks(
        doc,
        "Chapitre 4 : Réalisation et mise en œuvre",
        ["Chapitre 5"],
        normalize(ch4_blocks),
    )
    fill_with_blocks(
        doc,
        "Chapitre 5 : Tests, résultats, limites et perspectives",
        ["Conclusion générale"],
        normalize(ch5_blocks),
    )

    # Update conclusion générale to reflect realization + tests
    concl = find_para(doc, lambda t: t == "Conclusion générale", prefer_heading=True)
    if concl is not None:
        # replace following paragraphs until Bibliographie
        clear_until(doc, concl, ["Bibliographie", "Bibliographie et webographie"])
        anchor = concl
        for text in [
            "Ce projet est parti d'un écart concret : besoin d'évaluer la "
            "programmation sur machine, difficulté à garantir l'intégrité des "
            "rendus lorsque la génération de code devient banale. Cylentic propose "
            "une réponse pragmatique pour une salle d'examen d'école d'ingénieurs : "
            "une plateforme web multi établissement, un parcours étudiant contrôlé, "
            "une exécution isolée et un journal d'incidents lisible par le professeur.",
            "L'analyse des besoins a borné un MVP assumé. La conception UML a fixé "
            "les rôles, les données et les flux. La réalisation a ensuite livré une "
            "chaîne complète, de l'onboarding administrateur à la soumission "
            "étudiant, en passant par la publication, les contrôles de passation, "
            "la correction automatique et les exports. La campagne de validation a "
            "confirmé le comportement des parcours critiques après correction des "
            "anomalies de déploiement local, de sérialisation et de synchronisation "
            "du timer.",
            "Des limites demeurent, notamment face à un second appareil non "
            "contrôlé, face au multi langage et face à l'industrialisation du temps "
            "réel. Elles n'effacent pas le résultat principal : disposer d'un "
            "système utilisable dans son périmètre, justifié dans ses choix, et "
            "prêt pour des usages pilotes en salle. Sur le plan méthodologique, le "
            "travail a conduit l'équipe à tenir ensemble analyse, modélisation, "
            "implémentation et recette, sans perdre de vue le verdict d'un jury.",
        ]:
            anchor = insert_paragraph_after(anchor, text)

    # Update figure list with 4.1 and 4.2 if missing
    liste_fig = find_para(doc, lambda t: t.startswith("Liste des figures"))
    if liste_fig is not None:
        has41 = any("Figure 4.1" in p.text for p in doc.paragraphs)
        if not has41:
            # find last figure 3.26 line to insert after, else after heading
            anchor = liste_fig
            for p in doc.paragraphs:
                if "Figure 3.26" in p.text:
                    anchor = p
            anchor = insert_paragraph_after(
                anchor,
                "Figure 4.1 : Salle d'attente avant le démarrage de l'examen",
                align="left",
                first_line=False,
                space_after=3,
            )
            insert_paragraph_after(
                anchor,
                "Figure 4.2 : Environnement de composition pendant l'épreuve",
                align="left",
                first_line=False,
                space_after=3,
            )

    full = "\n".join(p.text for p in doc.paragraphs)
    if "—" in full:
        raise SystemExit("Em dash detected")
    if any(
        p.text.startswith("[Contenu du point 4.") or p.text.startswith("[Contenu du point 5.")
        for p in doc.paragraphs
    ):
        raise SystemExit("Placeholders remain")
    if "Personne C" in full or "fascicule" in full.lower():
        raise SystemExit("Meta person C leaked")

    doc.save(FINAL)
    print("Saved", FINAL)
    print("has 4.4.4", any("4.4.4" in p.text for p in doc.paragraphs))
    print("has 5.2", any(p.text.startswith("5.2 ") for p in doc.paragraphs))
    print("exports fixed", "exports de résultats déjà disponibles" in full or "Excel, PDF" in full)
    print("SADM", "SADM-" in full)
    print("Judge0 only état art-ish", full.count("Judge0"))


if __name__ == "__main__":
    main()
