#!/usr/bin/env python3
"""Génère le rapport Cylentic : carcasse complète + chapitre 3 final."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def set_run_font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_paragraph(
    doc,
    text="",
    *,
    size=12,
    bold=False,
    italic=False,
    align="justify",
    space_after=8,
    space_before=0,
    first_line=True,
):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line and align == "justify" and text:
        pf.first_line_indent = Cm(0.75)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_heading_styled(doc, text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(14 if level == 1 else 10)
    pf.space_after = Pt(10)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.keep_with_next = True
    size = 14 if level <= 2 else 12
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True)
    # Outline level for Word TOC
    p.style = doc.styles["Heading %d" % min(level, 3)]
    for r in p.runs:
        set_run_font(r, size=size, bold=True)
    return p


def page_break(doc):
    doc.add_page_break()


def add_placeholder_block(doc, label):
    add_paragraph(
        doc,
        label,
        italic=True,
        align="left",
        first_line=False,
        space_after=6,
    )


def add_figure_slot(doc, number, title, description):
    add_paragraph(
        doc,
        f"[Emplacement de la figure {number}]",
        italic=True,
        align="center",
        first_line=False,
        space_before=10,
        space_after=4,
    )
    add_paragraph(
        doc,
        description,
        italic=True,
        align="center",
        size=11,
        first_line=False,
        space_after=4,
    )
    add_paragraph(
        doc,
        f"Figure {number} : {title}",
        size=11,
        align="center",
        first_line=False,
        space_after=12,
    )


def add_table_title(doc, number, title):
    add_paragraph(
        doc,
        f"Tableau {number} : {title}",
        size=11,
        bold=True,
        align="left",
        first_line=False,
        space_before=8,
        space_after=4,
    )


def set_doc_defaults(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.5)  # 2.5 + reliure 1
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for i in range(1, 4):
        hs = doc.styles[f"Heading {i}"]
        hs.font.name = "Times New Roman"
        hs.font.color.rgb = None
        hs.font.bold = True
        hs.font.size = Pt(14 if i <= 2 else 12)


def build():
    doc = Document()
    set_doc_defaults(doc)

    # ========== PAGE DE GARDE ==========
    for _ in range(2):
        add_paragraph(doc, "", first_line=False, space_after=0)
    add_paragraph(
        doc,
        "BURKINA FASO",
        bold=True,
        align="center",
        first_line=False,
        space_after=2,
    )
    add_paragraph(
        doc,
        "La Patrie ou la Mort, nous Vaincrons",
        italic=True,
        align="center",
        first_line=False,
        space_after=10,
    )
    add_paragraph(
        doc,
        "Ministère de l'Enseignement Supérieur, de la Recherche et de l'Innovation",
        align="center",
        size=11,
        first_line=False,
        space_after=6,
    )
    add_paragraph(
        doc,
        "[LOGO DE L'ÉTABLISSEMENT]",
        italic=True,
        align="center",
        first_line=False,
        space_after=6,
    )
    add_paragraph(
        doc,
        "Nom de l'établissement",
        bold=True,
        align="center",
        first_line=False,
        space_after=4,
    )
    add_paragraph(
        doc,
        "Filière / Option : Informatique",
        align="center",
        first_line=False,
        space_after=14,
    )
    add_paragraph(
        doc,
        "RAPPORT DE PROJET DE FIN D'ANNÉE",
        bold=True,
        size=14,
        align="center",
        first_line=False,
        space_after=10,
    )
    add_paragraph(
        doc,
        "Cylentic : plateforme web sécurisée d'examens de programmation",
        bold=True,
        size=14,
        align="center",
        first_line=False,
        space_after=16,
    )
    add_paragraph(
        doc,
        "Présenté par :",
        align="center",
        first_line=False,
        space_after=4,
    )
    add_paragraph(doc, "Nom et prénoms de l'étudiant 1", align="center", first_line=False, space_after=2)
    add_paragraph(doc, "Nom et prénoms de l'étudiant 2", align="center", first_line=False, space_after=2)
    add_paragraph(doc, "Nom et prénoms de l'étudiant 3", align="center", first_line=False, space_after=12)
    add_paragraph(
        doc,
        "Encadrant principal : Nom, prénoms, grade et fonction",
        align="center",
        first_line=False,
        space_after=4,
    )
    add_paragraph(
        doc,
        "Encadrant : Nom, prénoms, grade et fonction",
        align="center",
        first_line=False,
        space_after=16,
    )
    add_paragraph(
        doc,
        "Année académique 20..-20..",
        align="center",
        first_line=False,
        space_after=4,
    )

    # ========== PAGE DE TITRE ==========
    page_break(doc)
    add_paragraph(
        doc,
        "Page de titre",
        bold=True,
        align="center",
        first_line=False,
        space_after=10,
    )
    add_paragraph(
        doc,
        "Cette page reprend, en noir et blanc, les informations de la page de garde "
        "(établissement, filière, titre du projet, auteurs, encadrants, année académique).",
        italic=True,
        first_line=False,
    )

    # ========== DEDICACE ==========
    page_break(doc)
    add_heading_styled(doc, "Dédicace", 1)
    add_paragraph(
        doc,
        "À nos familles, pour leur soutien discret et constant tout au long de cette année. "
        "À nos enseignants, qui nous ont transmis les bases nécessaires pour mener un projet "
        "de bout en bout. À tous ceux qui, de près ou de loin, ont cru à Cylentic avant même "
        "que la première ligne de code n'existe.",
    )
    add_paragraph(
        doc,
        "(Texte à personnaliser par l'équipe.)",
        italic=True,
        first_line=False,
    )

    # ========== REMERCIEMENTS ==========
    page_break(doc)
    add_heading_styled(doc, "Remerciements", 1)
    add_paragraph(
        doc,
        "Nous adressons nos remerciements à notre encadrant principal pour sa disponibilité, "
        "ses remarques précises et la rigueur qu'il a exigée à chaque étape du projet. "
        "Nous remercions également notre second encadrant pour ses conseils méthodologiques "
        "et techniques.",
    )
    add_paragraph(
        doc,
        "Nos remerciements vont à l'administration et au corps enseignant de notre établissement "
        "pour le cadre académique offert, ainsi qu'à nos camarades qui ont accepté de tester "
        "les premiers parcours de passation et de signaler les points durs de l'interface.",
    )
    add_paragraph(
        doc,
        "Enfin, nous remercions nos familles pour la patience et le soutien quotidien sans "
        "lesquels ce travail n'aurait pas abouti dans les délais impartis.",
    )
    add_paragraph(
        doc,
        "(Compléter ou modifier selon les personnes réellement à remercier.)",
        italic=True,
        first_line=False,
    )

    # ========== SOMMAIRE ==========
    page_break(doc)
    add_heading_styled(doc, "Sommaire", 1)
    sommaire = [
        "Dédicace",
        "Remerciements",
        "Liste des sigles et abréviations",
        "Liste des figures",
        "Liste des tableaux",
        "Résumé",
        "Abstract",
        "Introduction générale",
        "Chapitre 1 : Contexte, problématique et étude de l'existant",
        "Chapitre 2 : Analyse des besoins et spécification du système",
        "Chapitre 3 : Conception du système Cylentic",
        "Chapitre 4 : Réalisation et mise en œuvre",
        "Chapitre 5 : Tests, résultats, limites et perspectives",
        "Conclusion générale",
        "Bibliographie et webographie",
        "Annexes",
        "Table des matières",
    ]
    for item in sommaire:
        add_paragraph(doc, item, align="left", first_line=False, space_after=4)

    # ========== SIGLES ==========
    page_break(doc)
    add_heading_styled(doc, "Liste des sigles et abréviations", 1)
    sigles = [
        ("API", "Application Programming Interface"),
        ("CSV", "Comma Separated Values"),
        ("JWT", "JSON Web Token"),
        ("MVP", "Minimum Viable Product"),
        ("QCM", "Questionnaire à choix multiples"),
        ("REST", "Representational State Transfer"),
        ("SGBD", "Système de gestion de bases de données"),
        ("UML", "Unified Modeling Language"),
    ]
    for s, dfn in sigles:
        add_paragraph(doc, f"{s} : {dfn}", align="left", first_line=False, space_after=3)

    # ========== LISTE FIGURES ==========
    page_break(doc)
    add_heading_styled(doc, "Liste des figures", 1)
    figures_index = [
        "Figure 3.1 : Architecture logique de Cylentic",
        "Figure 3.2 : Déploiement de l'infrastructure",
        "Figure 3.3 : Cas d'utilisation du Super Admin",
        "Figure 3.4 : Cas d'utilisation de l'administrateur (organisation)",
        "Figure 3.5 : Cas d'utilisation de l'administrateur (comptes)",
        "Figure 3.6 : Cas d'utilisation du professeur (examens)",
        "Figure 3.7 : Cas d'utilisation du professeur (suivi)",
        "Figure 3.8 : Cas d'utilisation de l'étudiant",
        "Figure 3.9 : Classes organisation et comptes",
        "Figure 3.10 : Classes conception d'examen",
        "Figure 3.11 : Classes passation et correction",
        "Figure 3.12 : Objets organisation",
        "Figure 3.13 : Objets conception d'examen",
        "Figure 3.14 : Objets passation",
        "Figure 3.15 : Composants vue d'ensemble",
        "Figure 3.16 : Composants administration et authentification",
        "Figure 3.17 : Composants examens et passation",
        "Figure 3.18 : États du cycle de vie d'un examen",
        "Figure 3.19 : Séquence de connexion",
        "Figure 3.20 : Séquence de publication d'un examen",
        "Figure 3.21 : Séquence d'accès étudiant",
        "Figure 3.22 : Séquence d'exécution du code",
        "Figure 3.23 : Séquence de soumission et correction",
        "Figure 3.24 : Séquence d'enregistrement d'un incident",
        "Figure 3.25 : Activité accès et salle d'attente",
        "Figure 3.26 : Activité composition et soumission",
    ]
    for f in figures_index:
        add_paragraph(doc, f, align="left", first_line=False, space_after=3)

    # ========== LISTE TABLEAUX ==========
    page_break(doc)
    add_heading_styled(doc, "Liste des tableaux", 1)
    for t in [
        "Tableau 3.1 : Correspondance entre acteurs et responsabilités",
        "Tableau 3.2 : Principales tables persistantes et rôle",
    ]:
        add_paragraph(doc, t, align="left", first_line=False, space_after=3)

    # ========== RESUME ==========
    page_break(doc)
    add_heading_styled(doc, "Résumé", 1)
    add_paragraph(
        doc,
        "Les examens de programmation sur machine sont devenus difficiles à trancher : "
        "un étudiant peut obtenir une solution crédible en peu de temps grâce aux assistants "
        "de génération de code, tandis qu'un retour au papier élude la vraie compétence à "
        "développer et à exécuter. Dans le contexte d'une école d'ingénieurs au Burkina Faso, "
        "où les épreuves se déroulent souvent en salle informatique avec des contraintes "
        "matérielles réelles, nous proposons Cylentic, une plateforme web multi établissement "
        "d'examens sécurisés.",
    )
    add_paragraph(
        doc,
        "Le projet couvre la gestion des établissements et des comptes, la création et la "
        "publication d'examens, la passation contrôlée côté navigateur, l'exécution isolée "
        "du code Python, la journalisation des incidents et une correction automatique. "
        "Le présent rapport décrit la démarche, la conception UML, la réalisation du MVP "
        "et les premiers résultats d'évaluation.",
    )
    add_paragraph(
        doc,
        "Mots clés : examen de programmation, sécurité navigateur, sandbox, UML, Next.js, correction automatique.",
        first_line=False,
    )

    # ========== ABSTRACT ==========
    page_break(doc)
    add_heading_styled(doc, "Abstract", 1)
    add_paragraph(
        doc,
        "Programming exams on shared lab computers are hard to assess fairly when students "
        "can obtain working solutions in seconds with code generation tools. Paper exams "
        "avoid this issue but poorly measure the ability to write and run code. Cylentic "
        "is a multi tenant web platform designed for engineering schools in Burkina Faso. "
        "It combines browser based session controls, an isolated Python sandbox, incident "
        "logging and automated grading. This report presents the problem analysis, the UML "
        "design, the MVP implementation and early evaluation results.",
    )
    add_paragraph(
        doc,
        "Keywords: programming exam, browser security, sandbox, UML, Next.js, automated grading.",
        first_line=False,
    )

    # ========== INTRODUCTION ==========
    page_break(doc)
    add_heading_styled(doc, "Introduction générale", 1)
    add_paragraph(
        doc,
        "Former un informaticien, c'est aussi vérifier qu'il sait résoudre un problème "
        "sur machine. Or la nature des rendus a changé. Les assistants de génération de "
        "code produisent rapidement des réponses plausibles. Un devoir réalisé hors salle "
        "devient alors ambigu : le professeur voit un résultat, rarement le chemin qui y "
        "a mené. Revenir au papier protège contre certaines aides externes, mais il "
        "écarte l'exécution réelle, les erreurs de compilation et la logique d'un "
        "environnement de développement.",
    )
    add_paragraph(
        doc,
        "Dans nos écoles, l'épreuve pratique a généralement lieu en salle informatique. "
        "Les postes sont partagés, le réseau est parfois instable, et le surveillant "
        "reste présent. Ce cadre impose une solution légère côté étudiant, sans "
        "installation lourde, tout en donnant au professeur des éléments exploitables "
        "après l'épreuve : copie, journal d'incidents, scores issus des tests.",
    )
    add_paragraph(
        doc,
        "La question qui structure ce travail est simple à énoncer et exigeante à traiter : "
        "comment organiser un examen de programmation sur navigateur qui réduise "
        "fortement les possibilités de triche, conserve un rôle clair au surveillant "
        "humain, et restitue une correction lisible ? Cylentic répond à cette question "
        "sous la forme d'une plateforme multi établissement, avec des rôles distincts "
        "(super administrateur, administrateur d'établissement, professeur, étudiant), "
        "un parcours de passation verrouillé et une exécution de code isolée.",
    )
    add_paragraph(
        doc,
        "L'objectif général consiste à concevoir et réaliser un MVP opérationnel. "
        "Les objectifs spécifiques portent sur la gestion organisationnelle, la "
        "publication d'examens, le contrôle de la session navigateur, l'exécution "
        "sandbox, la journalisation des incidents et la correction automatique. "
        "La méthode suivie combine analyse des besoins, modélisation UML, "
        "implémentation sur une stack web moderne et tests de parcours.",
    )
    add_paragraph(
        doc,
        "Le rapport s'organise en cinq chapitres. Le premier situe le contexte et "
        "l'existant. Le deuxième formalise les besoins. Le troisième présente la "
        "conception, déjà développée dans ce document. Le quatrième décrit la "
        "réalisation. Le cinquième discute les tests, les limites et les perspectives. "
        "Une conclusion générale ferme le propos.",
    )

    # ========== CHAPITRE 1 ==========
    page_break(doc)
    add_heading_styled(doc, "Chapitre 1 : Contexte, problématique et étude de l'existant", 1)
    add_paragraph(
        doc,
        "Ce chapitre pose le cadre pédagogique et technique du projet, formule la "
        "problématique, précise les objectifs, compare les solutions existantes et "
        "positionne Cylentic.",
    )
    add_heading_styled(doc, "1.1 Contexte pédagogique et technologique", 2)
    add_placeholder_block(doc, "[Contenu du point 1.1 à coller ici]")
    add_heading_styled(doc, "1.2 Problématique", 2)
    add_placeholder_block(doc, "[Contenu du point 1.2 à coller ici]")
    add_heading_styled(doc, "1.3 Objectifs et hypothèse de travail", 2)
    add_placeholder_block(doc, "[Contenu du point 1.3 à coller ici]")
    add_heading_styled(doc, "1.4 État de l'art et comparaison des solutions existantes", 2)
    add_placeholder_block(doc, "[Contenu du point 1.4 à coller ici]")
    add_heading_styled(doc, "1.5 Positionnement de Cylentic", 2)
    add_placeholder_block(doc, "[Contenu du point 1.5 à coller ici]")
    add_heading_styled(doc, "1.6 Conclusion partielle", 2)
    add_placeholder_block(doc, "[Contenu du point 1.6 à coller ici]")

    # ========== CHAPITRE 2 ==========
    page_break(doc)
    add_heading_styled(doc, "Chapitre 2 : Analyse des besoins et spécification du système", 1)
    add_paragraph(
        doc,
        "Ce chapitre délimite le MVP, identifie les acteurs, exprime les besoins "
        "fonctionnels et non fonctionnels, rappelle les règles métier structurantes "
        "et présente les scénarios d'usage prioritaires.",
    )
    add_heading_styled(doc, "2.1 Périmètre du projet et du MVP", 2)
    add_placeholder_block(doc, "[Contenu du point 2.1 à coller ici]")
    add_heading_styled(doc, "2.2 Identification des acteurs", 2)
    add_placeholder_block(doc, "[Contenu du point 2.2 à coller ici]")
    add_heading_styled(doc, "2.3 Besoins fonctionnels", 2)
    add_placeholder_block(doc, "[Contenu du point 2.3 à coller ici]")
    add_heading_styled(doc, "2.4 Besoins non fonctionnels", 2)
    add_placeholder_block(doc, "[Contenu du point 2.4 à coller ici]")
    add_heading_styled(doc, "2.5 Règles métier structurantes", 2)
    add_placeholder_block(doc, "[Contenu du point 2.5 à coller ici]")
    add_heading_styled(doc, "2.6 Scénarios d'usage prioritaires", 2)
    add_placeholder_block(doc, "[Contenu du point 2.6 à coller ici]")
    add_heading_styled(doc, "2.7 Conclusion partielle", 2)
    add_placeholder_block(doc, "[Contenu du point 2.7 à coller ici]")

    # ========== CHAPITRE 3 COMPLET ==========
    page_break(doc)
    add_heading_styled(doc, "Chapitre 3 : Conception du système Cylentic", 1)
    add_paragraph(
        doc,
        "Une fois les besoins posés, il reste à donner au système une forme stable, "
        "discutable devant un jury et transposable en code. Ce chapitre décrit la "
        "démarche de conception, l'architecture retenue, les vues UML produites à "
        "partir du fonctionnement réel de la plateforme, la structure des données, "
        "puis les choix liés à la sécurité de passation et à la correction automatique.",
    )

    # 3.1
    add_heading_styled(doc, "3.1 Démarche de conception", 2)
    add_paragraph(
        doc,
        "Nous avons choisi UML comme langage commun pour raisonner sur Cylentic. "
        "Ce choix n'est pas purement scolaire. Les examens multi rôles, le verrouillage "
        "de session et la correction après soumission créent des enchaînements qu'il "
        "est difficile de tenir dans la seule tête d'un développeur. Les diagrammes "
        "servent à rendre ces enchaînements visibles, à vérifier les responsabilités "
        "et à éviter les contradictions entre l'interface, l'API et la base de données.",
    )
    add_paragraph(
        doc,
        "La démarche a été itérative. Les premiers modèles ont été confrontés au "
        "schéma Prisma, aux routes Next.js et aux services métier déjà présents dans "
        "le dépôt. Lorsque le code montrait une règle plus précise que le schéma "
        "initial, le modèle a été corrigé. Inversement, certains états encore partiellement "
        "automatisés (passage de publié à en cours, puis à terminé) ont été conservés "
        "dans la conception, car ils appartiennent au cycle de vie prévu de l'examen.",
    )
    add_paragraph(
        doc,
        "Les vues retenues couvrent les cas d'utilisation, les classes et objets, "
        "les composants, le déploiement, les états, les séquences et les activités. "
        "Chaque vue répond à une question différente. Les cas d'utilisation disent "
        "qui fait quoi. Les classes disent quelles informations persistent. Les "
        "séquences disent comment les messages circulent. Les activités disent "
        "comment un étudiant traverse la journée d'épreuve.",
    )
    add_paragraph(
        doc,
        "Nous avons aussi refusé deux dérives fréquentes. La première consiste à "
        "produire des diagrammes décoratifs, trop génériques pour guider le code. "
        "La seconde consiste à coller le schéma de base de données tel quel et à "
        "l'appeler conception. Ici, chaque figure a un rôle dans le raisonnement : "
        "expliquer une responsabilité, une contrainte ou un enchaînement critique "
        "du MVP.",
    )

    # 3.2
    add_heading_styled(doc, "3.2 Architecture générale et déploiement", 2)
    add_paragraph(
        doc,
        "Cylentic est une application web centrée sur Next.js. Le même processus "
        "sert les pages et les routes API. Cette concentration convient au MVP : "
        "moins de services à synchroniser le jour de l'examen, un déploiement plus "
        "simple sur un serveur unique. Les données durables sont confiées à MySQL. "
        "Redis est prévu pour les files et les besoins de cache liés aux workers. "
        "L'exécution du code étudiant passe par un service sandbox qui privilégie "
        "un conteneur Docker Python isolé, avec repli local si nécessaire.",
    )
    add_paragraph(
        doc,
        "Côté client, l'étudiant compose dans le navigateur avec l'éditeur Monaco. "
        "Les contrôles de passation (plein écran, visibilité d'onglet, garde du "
        "presse papiers, verrouillage clavier) vivent dans des hooks dédiés. "
        "Côté serveur, la logique métier est séparée des routes : authentification, "
        "examens, participation, correction et sandbox forment des services "
        "distincts. Deux workers complètent le tableau : un pour la correction "
        "en file, un pour les transitions de statut d'examen.",
    )
    add_paragraph(
        doc,
        "Le choix d'une application monolithique modulaire pour le MVP n'interdit "
        "pas une évolution ultérieure. Il fixe surtout une priorité : le jour J, "
        "un professeur doit pouvoir publier un examen, un étudiant doit pouvoir "
        "composer, et une copie doit pouvoir être corrigée sans dépendre d'une "
        "nuée de microservices encore fragiles. Les frontières internes "
        "(pages, API, services, persistance) préparent néanmoins une séparation "
        "plus nette si la charge ou l'équipe grandit.",
    )
    add_figure_slot(
        doc,
        "3.1",
        "Architecture logique de Cylentic",
        "Schéma à insérer : navigateur (pages et Monaco), application Next.js "
        "(middleware, API, services, workers), MySQL, Redis et sandbox Docker Python.",
    )
    add_paragraph(
        doc,
        "Le diagramme de déploiement traduit cette architecture en nœuds "
        "d'exécution. Le poste utilisateur héberge le navigateur. Le serveur "
        "d'application héberge Next.js et les workers. MySQL et Redis apparaissent "
        "comme nœuds de données. L'hôte d'exécution isolée accueille le conteneur "
        "sandbox. En production, le lien navigateur application est prévu en HTTPS. "
        "En développement local, le même produit tourne souvent en HTTP sur "
        "localhost ; le modèle de déploiement décrit la cible d'exploitation.",
    )
    add_figure_slot(
        doc,
        "3.2",
        "Déploiement de l'infrastructure",
        "Insérer le diagramme PlantUML 08_deploiement.puml (poste utilisateur, "
        "serveur d'application, MySQL, Redis, hôte sandbox).",
    )

    # 3.3
    add_heading_styled(doc, "3.3 Vue des cas d'utilisation", 2)
    add_paragraph(
        doc,
        "Les acteurs métier de Cylentic sont le Super Admin, l'administrateur "
        "d'établissement, le professeur et l'étudiant. Le surveillant physique "
        "intervient en salle sans compte applicatif. Cette répartition évite de "
        "mélanger la gouvernance de la plateforme, la gestion locale d'une école, "
        "la pédagogie de l'épreuve et la composition.",
    )
    add_table_title(doc, "3.1", "Correspondance entre acteurs et responsabilités")
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    cells = [
        ("Acteur", "Responsabilité principale"),
        ("Super Admin", "Administrer les établissements et le pilotage plateforme"),
        ("Administrateur d'établissement", "Organisation scolaire et comptes locaux"),
        ("Professeur", "Conception, publication et suivi des examens"),
        ("Étudiant", "Accès contrôlé, composition, soumission"),
    ]
    for i, (a, b) in enumerate(cells):
        table.rows[i].cells[0].text = a
        table.rows[i].cells[1].text = b
        for cell in table.rows[i].cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run_font(r, size=11, bold=(i == 0))
    add_paragraph(doc, "", first_line=False, space_after=8)

    add_paragraph(
        doc,
        "Le Super Admin opère hors du tenant scolaire. Il gère les établissements, "
        "les plans et le suivi global. Les cas d'utilisation correspondants sont "
        "regroupés dans une vue dédiée afin de ne pas noyer cette gouvernance dans "
        "le détail pédagogique.",
    )
    add_figure_slot(
        doc,
        "3.3",
        "Cas d'utilisation du Super Admin",
        "Insérer 01_usecase_super_admin.puml.",
    )
    add_paragraph(
        doc,
        "L'administrateur d'établissement structure l'espace local : années "
        "académiques, classes, import de listes, activation des comptes. Ses "
        "cas d'utilisation ont été scindés en deux diagrammes pour rester lisibles "
        "sur une page paysage : l'organisation d'abord, la gestion des comptes ensuite.",
    )
    add_figure_slot(
        doc,
        "3.4",
        "Cas d'utilisation de l'administrateur (organisation)",
        "Insérer 02a_usecase_admin_organisation.puml.",
    )
    add_figure_slot(
        doc,
        "3.5",
        "Cas d'utilisation de l'administrateur (comptes)",
        "Insérer 02b_usecase_admin_comptes.puml.",
    )
    add_paragraph(
        doc,
        "Le professeur porte la responsabilité pédagogique. Il crée les examens, "
        "compose les exercices de code et les QCM, publie l'épreuve (ce qui génère "
        "le code d'accès), suit la salle, consulte les résultats et lit le journal "
        "d'incidents. Là encore, deux vues séparent la conception de l'examen et "
        "le suivi pendant ou après la passation.",
    )
    add_figure_slot(
        doc,
        "3.6",
        "Cas d'utilisation du professeur (examens)",
        "Insérer 03a_usecase_professeur_examens.puml.",
    )
    add_figure_slot(
        doc,
        "3.7",
        "Cas d'utilisation du professeur (suivi)",
        "Insérer 03b_usecase_professeur_suivi.puml.",
    )
    add_paragraph(
        doc,
        "L'étudiant entre avec son identifiant, son mot de passe et le code d'examen. "
        "Il accepte les consignes, active le plein écran, attend en salle d'attente, "
        "compose, exécute son code pour tester, soumet, et peut être soumis "
        "automatiquement à l'expiration du temps ou après exclusion. Le diagramme "
        "étudiant concentre ce parcours sans le mélanger aux tâches d'administration.",
    )
    add_figure_slot(
        doc,
        "3.8",
        "Cas d'utilisation de l'étudiant",
        "Insérer 04_usecase_etudiant.puml.",
    )

    # 3.4
    add_heading_styled(doc, "3.4 Structure statique : classes et objets", 2)
    add_paragraph(
        doc,
        "Le modèle de classes découpe le domaine en trois paquets cohérents. "
        "Le premier couvre l'organisation et les comptes : administrateurs "
        "plateforme, établissements, plans d'abonnement, utilisateurs, profils "
        "professeur et étudiant, classes et années académiques. Le deuxième "
        "couvre la conception de l'examen : examen, rattachement aux classes, "
        "exercices, tests unitaires, questions et choix QCM. Le troisième couvre "
        "la passation : participation, sauvegardes automatiques, soumissions, "
        "résultats de tests, réponses QCM et incidents.",
    )
    add_paragraph(
        doc,
        "Les associations portent des noms explicites afin d'éviter une lecture "
        "ambiguë. Un administrateur plateforme administre des établissements. "
        "Un utilisateur a pour profil un enregistrement enseignant ou étudiant. "
        "Un examen rassemble des exercices. Une participation appartient à un "
        "étudiant pour un examen donné. Les clés étrangères du schéma Prisma "
        "apparaissent comme attributs afin de ne rien perdre lors du découpage "
        "en plusieurs figures.",
    )
    add_figure_slot(
        doc,
        "3.9",
        "Classes organisation et comptes",
        "Insérer 05a_classes_organisation.puml.",
    )
    add_figure_slot(
        doc,
        "3.10",
        "Classes conception d'examen",
        "Insérer 05b1_classes_conception_examen.puml.",
    )
    add_figure_slot(
        doc,
        "3.11",
        "Classes passation et correction",
        "Insérer 05b2_classes_passation.puml.",
    )
    add_paragraph(
        doc,
        "Les diagrammes d'objets illustrent le même découpage avec des instances "
        "concrètes. Ils facilitent la relecture pédagogique : un établissement "
        "exemple, un professeur nommé, un examen publié, une participation en "
        "cours, un incident de sortie du plein écran. Les attributs des objets "
        "reprennent l'intégralité des champs du modèle, sans abréviation.",
    )
    add_figure_slot(
        doc,
        "3.12",
        "Objets organisation",
        "Insérer 06a_objets_organisation.puml.",
    )
    add_figure_slot(
        doc,
        "3.13",
        "Objets conception d'examen",
        "Insérer 06b_objets_conception_examen.puml.",
    )
    add_figure_slot(
        doc,
        "3.14",
        "Objets passation",
        "Insérer 06c_objets_passation.puml.",
    )

    # 3.5
    add_heading_styled(doc, "3.5 Organisation en composants", 2)
    add_paragraph(
        doc,
        "Les composants décrivent le découpage logiciel sans retomber dans le "
        "détail des classes métier. La vue d'ensemble relie l'interface "
        "navigateur, le middleware d'authentification, la couche API, les "
        "services métier, l'accès aux données, les workers, MySQL, Redis et "
        "la sandbox. Deux vues complémentaires précisent d'une part "
        "l'administration et l'authentification, d'autre part le cœur examens "
        "et passation.",
    )
    add_figure_slot(
        doc,
        "3.15",
        "Composants vue d'ensemble",
        "Insérer 07a_composants_vue_ensemble.puml.",
    )
    add_figure_slot(
        doc,
        "3.16",
        "Composants administration et authentification",
        "Insérer 07b_composants_admin_auth.puml.",
    )
    add_figure_slot(
        doc,
        "3.17",
        "Composants examens et passation",
        "Insérer 07c_composants_examens.puml.",
    )
    add_paragraph(
        doc,
        "Cette organisation reflète une règle de conception suivie dans le dépôt : "
        "les pages ne portent pas la logique métier. Elles appellent l'API ; "
        "l'API s'appuie sur des services ; les services parlent à Prisma et, "
        "lorsque nécessaire, à la sandbox. Les workers restent proches des "
        "mêmes services pour éviter deux vérités divergentes sur la correction "
        "ou sur le statut d'un examen.",
    )

    # 3.6
    add_heading_styled(doc, "3.6 Dynamique du système", 2)
    add_heading_styled(doc, "3.6.1 Cycle de vie d'un examen", 3)
    add_paragraph(
        doc,
        "Un examen naît en brouillon. La publication le rend accessible via un "
        "code d'accès et fixe son statut à publié. Le démarrage le fait passer "
        "en cours. La clôture le termine. L'archivage le range hors du travail "
        "courant. La suppression reste possible tant que l'examen n'a pas "
        "basculé dans les états verrouillés de passation. Ce cycle, présent dans "
        "l'énumération ExamStatus, structure aussi le tableau de bord enseignant.",
    )
    add_figure_slot(
        doc,
        "3.18",
        "États du cycle de vie d'un examen",
        "Insérer 09_etats_examen.puml.",
    )

    add_heading_styled(doc, "3.6.2 Interactions principales", 3)
    add_paragraph(
        doc,
        "La connexion vérifie d'abord le format d'identifiant pour déduire le rôle. "
        "Elle compte les tentatives récentes, charge le compte (utilisateur ou "
        "administrateur plateforme), compare le mot de passe, écrit la tentative "
        "dans le journal, met à jour la dernière connexion et délivre un jeton "
        "placé en cookie de session. En cas d'échec répété, l'accès est refusé "
        "pour une durée de verrouillage.",
    )
    add_figure_slot(
        doc,
        "3.19",
        "Séquence de connexion",
        "Insérer 10_sequence_connexion.puml.",
    )
    add_paragraph(
        doc,
        "La publication est réservée au professeur propriétaire de l'examen. "
        "Le service contrôle que l'examen est encore modifiable et qu'il contient "
        "au moins un exercice, génère le code d'accès s'il n'existe pas encore, "
        "puis enregistre le statut publié. L'interface affiche ensuite le code "
        "à communiquer en salle.",
    )
    add_figure_slot(
        doc,
        "3.20",
        "Séquence de publication d'un examen",
        "Insérer 11_sequence_publication.puml.",
    )
    add_paragraph(
        doc,
        "L'accès étudiant enchaîne authentification avec code d'examen, page des "
        "consignes, activation du plein écran et création de la participation. "
        "Le service vérifie l'appartenance de l'étudiant à une classe autorisée "
        "et le respect du délai d'accès. La participation passe par les statuts "
        "connecté puis salle d'attente avant la composition.",
    )
    add_figure_slot(
        doc,
        "3.21",
        "Séquence d'accès étudiant",
        "Insérer 12_sequence_acces_etudiant.puml.",
    )
    add_paragraph(
        doc,
        "Pendant la composition, l'exécution d'essai envoie le code au service "
        "sandbox. Le chemin nominal lance un conteneur Docker sans réseau, avec "
        "limites mémoire et processeur. Si Docker n'est pas disponible, un "
        "repli local conserve la possibilité de tester. Le résultat revient à "
        "l'interface sous forme de sorties standard et d'erreur.",
    )
    add_figure_slot(
        doc,
        "3.22",
        "Séquence d'exécution du code",
        "Insérer 13_sequence_execution.puml.",
    )
    add_paragraph(
        doc,
        "La soumission, manuelle ou déclenchée par le timer, fige les copies à "
        "partir des sauvegardes automatiques et des réponses QCM, marque la "
        "participation comme terminée, puis appelle la correction. Pour chaque "
        "exercice de code en mode automatique, les tests unitaires sont exécutés "
        "dans la sandbox. Les scores partiels sont totalisés sur la participation. "
        "Le professeur conserve la possibilité d'ajuster ensuite la note finale.",
    )
    add_figure_slot(
        doc,
        "3.23",
        "Séquence de soumission et correction",
        "Insérer 14_sequence_soumission_correction.puml.",
    )
    add_paragraph(
        doc,
        "Les incidents (sortie du plein écran, changement d'onglet, collage, "
        "fermeture de session, souci réseau) sont enregistrés pendant la "
        "passation. Pour les sorties du plein écran, un seuil de deux occurrences "
        "entraîne l'exclusion : la participation est soumise avec le motif "
        "correspondant, puis corrigée. Cette règle est assumée comme garde "
        "technique. Elle ne remplace pas l'appréciation disciplinaire humaine.",
    )
    add_figure_slot(
        doc,
        "3.24",
        "Séquence d'enregistrement d'un incident",
        "Insérer 15_sequence_incidents.puml.",
    )

    add_heading_styled(doc, "3.6.3 Parcours d'activité de l'étudiant", 3)
    add_paragraph(
        doc,
        "Le diagramme d'activité raconte la journée d'épreuve sans se perdre "
        "dans les appels de méthodes. Il a été découpé en deux figures. La "
        "première couvre l'authentification, les consignes, le plein écran, "
        "la création de participation et l'attente jusqu'à l'heure de début. "
        "La seconde couvre la composition, les sauvegardes, les exécutions "
        "d'essai, le traitement des incidents, la soumission et la correction.",
    )
    add_figure_slot(
        doc,
        "3.25",
        "Activité accès et salle d'attente",
        "Insérer 16a_activite_acces_attente.puml.",
    )
    add_figure_slot(
        doc,
        "3.26",
        "Activité composition et soumission",
        "Insérer 16b_activite_composition_soumission.puml.",
    )

    # 3.7
    add_heading_styled(doc, "3.7 Conception des données", 2)
    add_paragraph(
        doc,
        "La persistance s'appuie sur MySQL via Prisma. Le schéma reprend les "
        "énumérations du domaine (rôles, statuts d'examen, modes de correction, "
        "types d'incident, raisons de soumission). Les tables centrales sont "
        "les établissements, les utilisateurs et profils, les classes, les "
        "examens, les exercices, les participations, les soumissions, les "
        "incidents et les journaux de tentatives de connexion ou de code.",
    )
    add_table_title(doc, "3.2", "Principales tables persistantes et rôle")
    table2 = doc.add_table(rows=9, cols=2)
    table2.style = "Table Grid"
    rows2 = [
        ("Table / ensemble", "Rôle"),
        ("establishments, subscription_plans", "Tenant et offre"),
        ("platform_admins", "Gouvernance plateforme"),
        ("users, profils, classes", "Organisation scolaire"),
        ("exams, exercises, unit_tests, qcm_*", "Conception pédagogique"),
        ("exam_participations", "Présence et statut de passation"),
        ("code_autosaves, submissions", "Contenu de copie"),
        ("incidents, login_attempts, exam_code_attempts", "Preuves et traçabilité"),
        ("submission_test_results", "Détail de correction automatique"),
    ]
    for i, (a, b) in enumerate(rows2):
        table2.rows[i].cells[0].text = a
        table2.rows[i].cells[1].text = b
        for cell in table2.rows[i].cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run_font(r, size=11, bold=(i == 0))
    add_paragraph(doc, "", first_line=False, space_after=8)
    add_paragraph(
        doc,
        "MySQL conserve la mémoire durable : sujets, copies, scores, journaux. "
        "Redis est positionné pour les besoins de file et de coordination des "
        "workers. Cette séparation évite de faire porter à la base relationnelle "
        "des responsabilités de file d'attente pour lesquelles elle est moins "
        "naturelle. En retour, la correction et le suivi pédagogique restent "
        "interrogeables par requêtes classiques après l'épreuve.",
    )
    add_paragraph(
        doc,
        "Plusieurs contraintes d'intégrité méritent d'être soulignées. Une "
        "participation est unique pour un couple examen étudiant. Les "
        "sauvegardes automatiques sont indexées par participation et exercice. "
        "Les soumissions suivent la même unicité. Les tentatives de code "
        "d'examen et de connexion sont historisées même en cas d'échec, car "
        "elles éclairent les fraudes répétées ou les erreurs de saisie le jour "
        "de l'épreuve. Ces choix relèvent autant de la conception que de "
        "l'exploitation pédagogique.",
    )

    # 3.8
    add_heading_styled(doc, "3.8 Conception de la sécurité de passation", 2)
    add_paragraph(
        doc,
        "La sécurité de Cylentic ne prétend pas supprimer toute fraude possible. "
        "Elle cherche à rendre coûteux et visible le contournement pendant une "
        "épreuve en salle. Le parcours commence par l'acceptation explicite des "
        "consignes et l'entrée en plein écran. Pendant la composition, les hooks "
        "détectent la perte de plein écran, le changement d'onglet, certains "
        "collages et des raccourcis sensibles. Chaque événement pertinent crée "
        "un incident horodaté rattaché à la participation.",
    )
    add_paragraph(
        doc,
        "Le seuil de deux sorties du plein écran déclenche une exclusion "
        "technique avec soumission forcée. Un avertissement précède l'exclusion "
        "et exige le retour en plein écran. Les autres types d'incidents "
        "alimentent le journal consultable par le professeur. Cette asymétrie "
        "est volontaire : tout signaler, mais n'automatiser l'exclusion que sur "
        "un critère simple, mesurable et expliqué à l'étudiant.",
    )
    add_paragraph(
        doc,
        "Deux principes bornent le dispositif. Premier principe : le surveillant "
        "humain reste nécessaire. Un second appareil, hors navigateur surveillé, "
        "échappe largement aux hooks. Second principe : le professeur lit le "
        "journal. La plateforme fournit des preuves ; elle ne rend pas un "
        "jugement disciplinaire à la place de l'institution.",
    )
    add_paragraph(
        doc,
        "Côté serveur, l'authentification s'appuie sur des jetons JWT en cookie "
        "httpOnly, un middleware de contrôle des espaces par rôle, une "
        "limitation des tentatives de connexion et un contrôle du code d'examen. "
        "Ces mécanismes complètent la surveillance navigateur sans s'y "
        "substituer.",
    )
    add_paragraph(
        doc,
        "Dans une salle burkinabè typique, cette conception se lit concrètement. "
        "Le surveillant circule, le navigateur enferme la session, le serveur "
        "garde l'heure et les preuves. Si le réseau vacille brièvement, les "
        "sauvegardes automatiques et la reprise de participation réduisent la "
        "perte de travail. Le système ne promet pas une invulnérabilité. Il "
        "promet une passation plus traçable qu'une session libre sur un poste "
        "partagé.",
    )

    # 3.9
    add_heading_styled(doc, "3.9 Conception de la correction automatique", 2)
    add_paragraph(
        doc,
        "La correction automatique repose sur des tests unitaires associés aux "
        "exercices de code. Après soumission, le service de correction exécute "
        "le code de l'étudiant pour chaque test dans la sandbox, compare la "
        "sortie normalisée à la sortie attendue, enregistre le résultat et "
        "calcule un score pondéré. Les QCM sont corrigés par comparaison des "
        "choix retenus aux choix corrects, selon le type de réponse (unique ou "
        "multiple).",
    )
    add_paragraph(
        doc,
        "Deux modes coexistent. En mode automatique, le score calculé alimente "
        "la note proposée. En mode manuel, l'exercice attend l'appréciation "
        "du professeur. Dans tous les cas, la note finale reste ajustable. "
        "Cette conception évite de transformer la plateforme en boîte noire "
        "où un barème machine écraserait le jugement enseignant.",
    )
    add_paragraph(
        doc,
        "Le lien avec la sandbox est le même que pour l'exécution d'essai "
        "pendant l'épreuve. La différence tient à l'intention : l'essai montre "
        "une sortie à l'étudiant ; la correction compare une sortie à une "
        "référence privée et produit des preuves stockées (statut du test, "
        "durée, éventuelle erreur).",
    )

    # 3.10
    add_heading_styled(doc, "3.10 Conclusion partielle", 2)
    add_paragraph(
        doc,
        "La conception de Cylentic articule une architecture web unique autour "
        "de Next.js, une modélisation UML alignée sur le domaine réel, un "
        "déploiement clair, un cycle de vie d'examen explicite et des flux "
        "dynamiques pour les moments sensibles de l'épreuve. Les choix de "
        "sécurité et de correction assument un MVP : des gardes techniques "
        "lisibles, un journal exploitable, une note machine qui n'efface pas "
        "le professeur. Sur cette base, la réalisation décrite au chapitre "
        "suivant peut s'évaluer sans ambiguïté fonctionnelle.",
    )

    # ========== CHAPITRE 4 ==========
    page_break(doc)
    add_heading_styled(doc, "Chapitre 4 : Réalisation et mise en œuvre", 1)
    add_paragraph(
        doc,
        "Ce chapitre décrit l'environnement de travail, les choix technologiques, "
        "l'organisation du code, la réalisation des modules principaux, les points "
        "techniques difficiles et le déploiement du MVP.",
    )
    add_heading_styled(doc, "4.1 Environnement de travail", 2)
    add_placeholder_block(doc, "[Contenu du point 4.1 à coller ici]")
    add_heading_styled(doc, "4.2 Choix technologiques", 2)
    add_placeholder_block(doc, "[Contenu du point 4.2 à coller ici]")
    add_heading_styled(doc, "4.3 Organisation du code et architecture applicative", 2)
    add_placeholder_block(doc, "[Contenu du point 4.3 à coller ici]")
    add_heading_styled(doc, "4.4 Réalisation des modules principaux", 2)
    add_placeholder_block(doc, "[Contenu du point 4.4 à coller ici]")
    add_heading_styled(doc, "4.5 Points techniques difficiles", 2)
    add_placeholder_block(doc, "[Contenu du point 4.5 à coller ici]")
    add_heading_styled(doc, "4.6 Déploiement", 2)
    add_placeholder_block(doc, "[Contenu du point 4.6 à coller ici]")
    add_heading_styled(doc, "4.7 Conclusion partielle", 2)
    add_placeholder_block(doc, "[Contenu du point 4.7 à coller ici]")

    # ========== CHAPITRE 5 ==========
    page_break(doc)
    add_heading_styled(doc, "Chapitre 5 : Tests, résultats, limites et perspectives", 1)
    add_paragraph(
        doc,
        "Ce chapitre présente la stratégie de tests, les jeux de résultats, la "
        "discussion des apports observés, les limites résiduelles et les "
        "perspectives de poursuite.",
    )
    add_heading_styled(doc, "5.1 Stratégie de tests", 2)
    add_placeholder_block(doc, "[Contenu du point 5.1 à coller ici]")
    add_heading_styled(doc, "5.2 Jeux de tests et résultats", 2)
    add_placeholder_block(doc, "[Contenu du point 5.2 à coller ici]")
    add_heading_styled(doc, "5.3 Discussion des résultats", 2)
    add_placeholder_block(doc, "[Contenu du point 5.3 à coller ici]")
    add_heading_styled(doc, "5.4 Limites et risques résiduels", 2)
    add_placeholder_block(doc, "[Contenu du point 5.4 à coller ici]")
    add_heading_styled(doc, "5.5 Perspectives", 2)
    add_placeholder_block(doc, "[Contenu du point 5.5 à coller ici]")
    add_heading_styled(doc, "5.6 Conclusion partielle", 2)
    add_placeholder_block(doc, "[Contenu du point 5.6 à coller ici]")

    # ========== CONCLUSION ==========
    page_break(doc)
    add_heading_styled(doc, "Conclusion générale", 1)
    add_paragraph(
        doc,
        "Ce projet est parti d'un écart concret : besoin d'évaluer la "
        "programmation sur machine, difficulté à garantir l'intégrité des "
        "rendus lorsque la génération de code devient banale. Cylentic propose "
        "une réponse pragmatique pour une salle d'examen d'école d'ingénieurs : "
        "une plateforme web multi établissement, un parcours étudiant contrôlé, "
        "une exécution isolée et un journal d'incidents lisible par le professeur.",
    )
    add_paragraph(
        doc,
        "La conception détaillée au chapitre 3 a permis de figer les rôles, les "
        "données et les flux avant d'entrer dans la réalisation. Le MVP vise "
        "l'essentiel plutôt que l'exhaustif. Des limites demeurent, notamment "
        "face à un second appareil non contrôlé et face à la montée en charge "
        "des exécutions simultanées. Ces limites, comme les perspectives de "
        "langages supplémentaires ou d'exports, seront précisées après les "
        "tests rapportés au chapitre 5.",
    )
    add_paragraph(
        doc,
        "Sur le plan méthodologique, le travail a conduit l'équipe à tenir "
        "ensemble analyse, modélisation et implémentation, sans perdre de vue "
        "le verdict d'un jury : comprendre le problème, montrer des choix "
        "justifiés, livrer un système utilisable dans son périmètre.",
    )

    # ========== BIBLIOGRAPHIE ==========
    page_break(doc)
    add_heading_styled(doc, "Bibliographie et webographie", 1)
    biblio = [
        "[1] Object Management Group, OMG Unified Modeling Language (OMG UML), version courante, documentation officielle.",
        "[2] Documentation Next.js, https://nextjs.org/docs (page consultée lors de la rédaction du projet).",
        "[3] Documentation Prisma, https://www.prisma.io/docs (page consultée lors de la rédaction du projet).",
        "[4] Documentation MySQL 8, https://dev.mysql.com/doc/ (page consultée lors de la rédaction du projet).",
        "[5] Documentation Redis, https://redis.io/docs/ (page consultée lors de la rédaction du projet).",
        "[6] Docker Documentation, https://docs.docker.com/ (page consultée lors de la rédaction du projet).",
        "[7] Guide du stagiaire Master/Ingénieur, ESTA, 2024-2025 (référence de mise en forme).",
        "[8] Compléter ici les articles et mémoires réellement utilisés pour l'état de l'art du chapitre 1.",
    ]
    for b in biblio:
        add_paragraph(doc, b, align="left", first_line=False, space_after=6)

    # ========== ANNEXES ==========
    page_break(doc)
    add_heading_styled(doc, "Annexes", 1)
    add_heading_styled(doc, "Annexe A : Synthèse du cahier des charges", 2)
    add_placeholder_block(doc, "[Contenu à coller ou à compléter]")
    add_heading_styled(doc, "Annexe B : Dictionnaire de données", 2)
    add_placeholder_block(doc, "[Contenu à coller ou à compléter]")
    add_heading_styled(doc, "Annexe C : Diagrammes UML en grand format", 2)
    add_placeholder_block(doc, "[Contenu à coller ou à compléter]")
    add_heading_styled(doc, "Annexe D : Extraits de code utiles", 2)
    add_placeholder_block(doc, "[Contenu à coller ou à compléter]")
    add_heading_styled(doc, "Annexe E : Modèle CSV d'import des étudiants", 2)
    add_placeholder_block(doc, "[Contenu à coller ou à compléter]")
    add_heading_styled(doc, "Annexe F : Captures d'écran complémentaires", 2)
    add_placeholder_block(doc, "[Contenu à coller ou à compléter]")

    # ========== TABLE DES MATIERES ==========
    page_break(doc)
    add_heading_styled(doc, "Table des matières", 1)
    add_paragraph(
        doc,
        "Mettre à jour automatiquement dans Word (Références > Table des matières) "
        "après finalisation des contenus et de la pagination.",
        italic=True,
        first_line=False,
    )

    out = "/workspace/docs/Rapport_Cylentic_Carcasse_Chapitre3.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    build()
