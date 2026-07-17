#!/usr/bin/env python3
"""Intègre le travail corrigé de la personne A dans le rapport final."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

FINAL = "/workspace/docs/Rapport_Cylentic_Carcasse_Chapitre3.docx"


def set_run_font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def insert_paragraph_after(paragraph, text, **kwargs):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    run = new_para.add_run(text)
    size = kwargs.get("size", 12)
    bold = kwargs.get("bold", False)
    italic = kwargs.get("italic", False)
    align = kwargs.get("align", "justify")
    first_line = kwargs.get("first_line", True)
    space_after = kwargs.get("space_after", 8)
    set_run_font(run, size=size, bold=bold, italic=italic)
    if align == "center":
        new_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "left":
        new_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        new_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = new_para.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line and align == "justify":
        pf.first_line_indent = Cm(0.75)
    return new_para


def find_para(doc, pred):
    for p in doc.paragraphs:
        if pred(p.text.strip()):
            return p
    return None


def clear_until(doc, heading_p, stops):
    to_remove = []
    started = False
    for p in list(doc.paragraphs):
        if p._p is heading_p._p:
            started = True
            continue
        if not started:
            continue
        t = p.text.strip()
        if any(t.startswith(sp) for sp in stops):
            break
        to_remove.append(p)
    for p in to_remove:
        el = p._element
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)


def insert_comparison_table(doc, anchor):
    anchor = insert_paragraph_after(
        anchor,
        "Tableau 1.1 : Comparaison des solutions existantes et de Cylentic",
        size=11,
        bold=True,
        align="left",
        first_line=False,
        space_after=4,
    )
    table = doc.add_table(rows=8, cols=6)
    table.style = "Table Grid"
    data = [
        [
            "Critère",
            "freeCodeCamp",
            "Moodle",
            "Proctoring",
            "Sandbox (ex. Judge0)",
            "Cylentic",
        ],
        ["Exécution de code", "Oui", "Non", "Non", "Oui", "Oui"],
        ["Verrouillage de session", "Non", "Non", "Partiel", "Non", "Oui"],
        ["Journal d'incidents", "Non", "Partiel", "Oui", "Non", "Oui"],
        ["Correction automatique", "Partiel", "Oui (QCM)", "Non", "Non", "Oui"],
        ["Multi établissements", "Non", "Oui", "Oui", "Non", "Oui"],
        ["Auto hébergement possible", "Non", "Oui", "Non", "Oui", "Oui"],
        ["Résilience réseau", "Non", "Partiel", "Non", "Non", "Oui"],
    ]
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = val
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, size=10, bold=(i == 0))
    tbl = table._tbl
    tbl.getparent().remove(tbl)
    anchor._p.addnext(tbl)
    new_p = OxmlElement("w:p")
    tbl.addnext(new_p)
    return Paragraph(new_p, anchor._parent)


def fill_section(doc, heading, items, stops):
    heading_p = find_para(doc, lambda t, h=heading: t == h or t.startswith(h))
    if heading_p is None:
        raise SystemExit(f"Missing heading: {heading}")
    clear_until(doc, heading_p, stops)
    anchor = heading_p
    for text, opts in items:
        if text == "TABLEAU_COMPARAISON":
            anchor = insert_comparison_table(doc, anchor)
            continue
        anchor = insert_paragraph_after(anchor, text, **opts)


def main():
    doc = Document(FINAL)

    resume_paras = [
        (
            "L'enseignement de la programmation se heurte aujourd'hui à un problème concret : "
            "un étudiant peut obtenir une solution fonctionnelle en peu de temps grâce à un "
            "assistant de génération de code, sans que le professeur puisse le démontrer. "
            "Cylentic est une plateforme web multi établissement qui permet de composer un "
            "examen de programmation dans le navigateur, dans un cadre contrôlé, avec "
            "exécution isolée du code Python, journalisation des incidents et correction "
            "automatique par tests unitaires.",
            {},
        ),
        (
            "Le travail a d'abord porté sur l'existant (plateformes d'apprentissage, LMS, "
            "outils de proctoring, sandboxes d'exécution) afin d'identifier l'espace laissé "
            "libre, puis sur un cahier des charges borné à un MVP réaliste. Ce MVP couvre "
            "Python, la gestion des établissements et des comptes, la création et la "
            "publication d'examens, la salle d'attente, l'environnement de composition "
            "sécurisé et la correction automatique. Sur cette base, la conception UML et "
            "la réalisation technique peuvent s'enchaîner sans ambiguïté de périmètre.",
            {},
        ),
        (
            "Mots clés : examen de programmation, sécurité navigateur, correction "
            "automatique, sandbox, plateforme web, UML.",
            {"first_line": False},
        ),
    ]

    abstract_paras = [
        (
            "Teaching computer programming today faces a concrete problem: a student can "
            "obtain a working solution quickly with a code generation assistant, without "
            "the instructor being able to prove it. Cylentic is a multi tenant web platform "
            "for taking a programming exam in the browser under controlled conditions, with "
            "isolated Python execution, incident logging and automatic grading through unit "
            "tests.",
            {},
        ),
        (
            "The project first reviewed existing tools (learning platforms, LMS, proctoring "
            "products, execution sandboxes) to locate the remaining gap, then defined a "
            "realistic MVP. This MVP covers Python, multi institution management, exam "
            "creation and publication, a waiting room, a secured composition environment "
            "and automatic grading. This foundation supports the UML design and the "
            "technical implementation that follow.",
            {},
        ),
        (
            "Keywords: programming exam, browser security, automatic grading, sandbox, "
            "web platform, UML.",
            {"first_line": False},
        ),
    ]

    intro_paras = [
        (
            "L'enseignement de la programmation repose depuis longtemps sur des exercices "
            "et des devoirs que l'étudiant réalise hors de la salle. Cette pratique vise à "
            "vérifier qu'il sait construire une solution, la tester et la corriger. "
            "Aujourd'hui, cette vérification est fragilisée par la généralisation des "
            "assistants de génération de code. Un étudiant peut obtenir rapidement un "
            "programme plausible, sans avoir conçu ni maîtrisé la solution. Dans bien des "
            "cas, le professeur ne dispose d'aucun moyen fiable pour distinguer un travail "
            "réellement produit d'un travail fortement assisté.",
            {},
        ),
        (
            "Face à ce constat, certains établissements reviennent aux examens papier. "
            "Cela réduit une partie de l'aide externe, mais crée un décalage pédagogique : "
            "dans le métier, on n'écrit pas du code au stylo. On utilise un éditeur, on "
            "exécute, on lit les erreurs et on corrige. Évaluer surtout la syntaxe sur "
            "feuille ne mesure plus la compétence recherchée, à savoir résoudre un problème "
            "sur machine.",
            {},
        ),
        (
            "Ce travail s'inscrit dans une école d'ingénieurs au Burkina Faso, en filière "
            "informatique. Les examens s'y déroulent le plus souvent en salle informatique, "
            "sur des postes partagés, avec une qualité de connexion variable. Une solution "
            "adaptée doit rester simple à déployer, fonctionner dans le navigateur, sans "
            "installation lourde, et tolérer des coupures réseau ponctuelles.",
            {},
        ),
        (
            "La question structurante du rapport peut se formuler ainsi : comment permettre "
            "à un étudiant de composer un examen de programmation sur ordinateur, tout en "
            "réduisant fortement les possibilités de triche, et en donnant au professeur "
            "des preuves exploitables en cas d'incident ?",
            {},
        ),
        (
            "L'objectif général est de concevoir et de réaliser une plateforme web "
            "d'examens de programmation sécurisés, nommée Cylentic. Les objectifs "
            "spécifiques sont les suivants : permettre à un établissement de gérer ses "
            "comptes et ses classes, permettre à un professeur de créer et de publier un "
            "examen, enfermer l'étudiant dans un environnement de composition contrôlé, "
            "exécuter son code de façon isolée, corriger automatiquement les exercices à "
            "l'aide de tests unitaires, et produire un journal d'incidents exploitable.",
            {},
        ),
        (
            "La méthode a consisté à étudier d'abord les solutions existantes afin "
            "d'identifier ce qu'elles couvrent et ce qu'elles laissent de côté, puis à "
            "spécifier un cahier des charges limité à un MVP réaliste, avant de concevoir "
            "le système en UML et de le développer.",
            {},
        ),
        (
            "Le rapport s'organise en cinq chapitres. Le premier présente le contexte, la "
            "problématique, les objectifs et l'état de l'art. Le deuxième formalise les "
            "besoins et le cahier des charges. Le troisième développe la conception du "
            "système. Le quatrième décrit la réalisation. Le cinquième discute les tests, "
            "les limites et les perspectives. Une conclusion générale ferme le propos.",
            {},
        ),
    ]

    ch11 = [
        (
            "Dans une formation en informatique, les devoirs et exercices de programmation "
            "occupent une place centrale. Ils ne servent pas uniquement à vérifier qu'un "
            "étudiant connaît la syntaxe d'un langage : ils mesurent sa capacité à analyser "
            "un problème, à le décomposer, à écrire une solution, à la tester, à "
            "interpréter les erreurs d'exécution et à corriger son code jusqu'à un résultat "
            "correct. Cette compétence pratique ne peut être évaluée de façon fiable qu'en "
            "observant l'étudiant produire réellement du code sur une machine.",
            {},
        ),
        (
            "C'est précisément cette différence qui distingue une évaluation sur machine "
            "d'une évaluation sur papier. Réciter une syntaxe sur une feuille demande de "
            "mémoriser la forme d'une instruction, sans jamais vérifier que le programme "
            "s'exécute. Un étudiant peut ainsi réussir en connaissant des structures de "
            "code, sans être capable de faire fonctionner un vrai programme face à un "
            "interpréteur. L'examen papier évalue donc une compétence voisine, mais pas la "
            "même : il ne mesure ni la capacité à faire tourner du code, ni celle de lire "
            "un message d'erreur, ni celle de déboguer.",
            {},
        ),
        (
            "Ce contexte a été transformé ces dernières années par la démocratisation des "
            "outils de génération de code. Des assistants comme ChatGPT ou Copilot "
            "permettent de décrire un problème en langage naturel et d'obtenir rapidement "
            "un programme plausible. Ces outils ont changé la nature des devoirs faits à "
            "la maison : un rendu individuel de qualité ne garantit plus que l'étudiant "
            "maîtrise la compétence évaluée.",
            {},
        ),
        (
            "Le numérique ne remplace pas totalement la surveillance humaine, il la "
            "complète. Un surveillant présent en salle reste utile pour observer des "
            "comportements que la plateforme ne voit pas, comme l'usage d'un téléphone ou "
            "une communication entre étudiants. Cette idée, selon laquelle le numérique "
            "appuie l'humain sans s'y substituer, structure une grande partie des choix "
            "de Cylentic.",
            {},
        ),
    ]

    ch12 = [
        (
            "Le problème central peut être formulé nettement : comment permettre à un "
            "étudiant de composer un examen de programmation sur ordinateur, tout en "
            "réduisant fortement les possibilités de triche, et en donnant au professeur "
            "des preuves exploitables en cas d'incident ?",
            {},
        ),
        (
            "Ce problème se découpe en sous problèmes concrets. Le premier concerne le "
            "contrôle de l'environnement de composition : empêcher de sortir de la zone "
            "d'examen, d'ouvrir un autre onglet ou de coller du contenu préparé. Le "
            "deuxième concerne la détection des comportements suspects et une réaction "
            "proportionnée. Le troisième concerne la correction, qui doit rester fiable "
            "quand plusieurs étudiants soumettent au même moment. Le quatrième concerne "
            "la traçabilité : le professeur doit disposer d'un historique précis. Le "
            "cinquième concerne les coupures réseau, qui ne doivent pas pénaliser "
            "injustement un étudiant.",
            {},
        ),
        (
            "Dans le contexte burkinabè visé, les examens se déroulent en salle "
            "informatique, sur des machines partagées, avec une qualité de connexion "
            "variable. La solution recherchée doit donc être simple à déployer et à "
            "utiliser le jour J, sans configuration complexe, tout en restant tolérante "
            "aux incidents réseau ponctuels.",
            {},
        ),
    ]

    ch13 = [
        (
            "L'objectif général de ce projet est de concevoir et de réaliser une "
            "plateforme web d'examens de programmation sécurisés.",
            {},
        ),
        (
            "Les objectifs spécifiques sont les suivants : gérer les établissements et "
            "les comptes rattachés, permettre au professeur de créer et de publier un "
            "examen, enfermer l'étudiant dans un environnement de composition contrôlé "
            "pendant l'épreuve, exécuter le code de façon isolée, corriger "
            "automatiquement les exercices à l'aide de tests unitaires définis par le "
            "professeur, et produire un journal d'incidents exploitable après l'épreuve.",
            {},
        ),
        (
            "L'hypothèse de travail retenue est la suivante : si l'on verrouille "
            "suffisamment le navigateur de l'étudiant et que l'on journalise "
            "systématiquement les événements suspects, alors il devient possible de "
            "restaurer une crédibilité utile aux examens pratiques sur machine, sans "
            "revenir au papier.",
            {},
        ),
    ]

    ch14 = [
        (
            "Plusieurs catégories de solutions existent déjà et couvrent, chacune "
            "partiellement, une partie du problème.",
            {},
        ),
        (
            "Les plateformes d'apprentissage en ligne, comme freeCodeCamp, proposent une "
            "interface proche de celle recherchée : un énoncé à côté d'un éditeur, avec "
            "exécution immédiate. Elles sont conçues pour l'apprentissage libre, non pour "
            "l'examen surveillé. Elles ne disposent en général d'aucun verrouillage de "
            "session digne d'une épreuve, ni d'un journal d'incidents orienté professeur.",
            {},
        ),
        (
            "Les systèmes de gestion de l'apprentissage, dont Moodle est le représentant "
            "le plus répandu, permettent de créer des cours, de déposer des devoirs, de "
            "gérer des classes et des comptes, et de configurer des créneaux d'examen. "
            "Leur module de quiz autorise une correction automatique de questions à choix "
            "multiples. Moodle reste cependant généraliste : il ne propose pas nativement "
            "un éditeur de code avec exécution isolée, et la surveillance poussée passe "
            "souvent par des extensions tierces, parfois coûteuses et lourdes à déployer.",
            {},
        ),
        (
            "Les outils de proctoring du marché, fondés sur la webcam, la détection de "
            "comportements ou le verrouillage de poste, répondent surtout à la "
            "surveillance d'examens généralistes. Leur coût de licence est souvent élevé. "
            "Ils ne savent ni exécuter ni corriger du code de programmation dans le sens "
            "attendu ici.",
            {},
        ),
        (
            "Les sandboxes d'exécution, dont Judge0 est un exemple connu, résolvent un "
            "problème ciblé : exécuter du code dans un environnement isolé, sans accès "
            "réseau, avec des limites de mémoire et de temps, puis renvoyer la sortie. "
            "Ces composants sont utiles et souvent auto hébergeables. Ils ne proposent "
            "toutefois ni interface d'examen, ni gestion d'établissements, ni mécanismes "
            "de passation contrôlée : ce sont des briques techniques, pas une plateforme "
            "d'épreuve.",
            {},
        ),
        ("TABLEAU_COMPARAISON", {}),
        (
            "Ce tour d'horizon montre qu'aucune solution ne couvre l'ensemble du besoin. "
            "Les plateformes d'apprentissage et les LMS gèrent surtout les contenus et "
            "les comptes. Les outils de proctoring gèrent la surveillance mais pas "
            "l'exécution pédagogique du code. Les sandboxes gèrent l'exécution isolée et "
            "presque rien d'autre. C'est dans cet espace laissé libre que se positionne "
            "Cylentic.",
            {},
        ),
    ]

    ch15 = [
        (
            "Cylentic n'est pas un clone d'une plateforme d'apprentissage ouverte, ni une "
            "simple variante d'un LMS généraliste. Ces outils sont pensés pour un cadre "
            "ouvert où l'aide externe n'est pas le problème central. Cylentic répond à un "
            "besoin différent : l'évaluation surveillée de la programmation.",
            {},
        ),
        (
            "Ce qui le distingue, c'est le couplage entre des éléments rarement réunis "
            "pour une épreuve de code en salle : éditeur dans le navigateur, exécution "
            "isolée via une sandbox Docker Python, contrôles de passation côté client, "
            "timers et règles métier côté serveur, journal d'incidents pour le "
            "professeur, et maintien du rôle du surveillant physique. Cylentic ajoute "
            "aussi une gouvernance multi établissement, avec un Super Admin plateforme "
            "distinct de l'administrateur d'école.",
            {},
        ),
        (
            "C'est cette combinaison, plus que chaque brique prise seule, qui justifie "
            "le projet : offrir une plateforme conçue dès le départ pour l'examen de "
            "programmation surveillé, plutôt qu'un assemblage fragile d'outils "
            "incompatibles.",
            {},
        ),
    ]

    ch16 = [
        (
            "Ce chapitre a montré que le problème est réel : les assistants de génération "
            "de code fragilisent les devoirs traditionnels, sans que le retour au papier "
            "soit une réponse pédagogique satisfaisante. L'existant ne couvre le besoin "
            "que par fragments. Cylentic trouve sa place dans cet intervalle. Le chapitre "
            "suivant précise l'analyse des besoins et le cahier des charges.",
            {},
        ),
    ]

    ch21 = [
        (
            "Le projet assume un périmètre volontairement limité pour sa première "
            "version. Le MVP inclut : le langage Python, l'authentification avec rôles "
            "(Super Admin, administrateur d'établissement, professeur, étudiant), la "
            "création d'examens, la salle d'attente, l'environnement de composition "
            "sécurisé, les mécanismes de contrôle navigateur, la correction automatique "
            "par tests unitaires, un module QCM de base, et la gestion de plusieurs "
            "établissements.",
            {},
        ),
        (
            "Sont exclus de ce périmètre : le support d'autres langages que Python, la "
            "facturation réelle des établissements, l'export PDF ou Excel des résultats, "
            "la consultation en direct du code pendant l'épreuve, et l'activation de "
            "compte par email. Le MVP s'appuie plutôt sur des comptes créés par "
            "l'administration et, le cas échéant, un changement de mot de passe imposé.",
            {},
        ),
        (
            "Ce périmètre restreint concentre l'effort sur le cœur du problème : la "
            "passation contrôlée et la correction automatique, plutôt que sur des "
            "fonctions secondaires.",
            {},
        ),
    ]

    ch22 = [
        (
            "La plateforme distingue plusieurs acteurs, chacun avec un rôle précis.",
            {},
        ),
        (
            "Le Super Admin administre la plateforme : établissements, plans et pilotage "
            "global, hors du détail pédagogique d'une école.",
            {},
        ),
        (
            "L'administrateur d'établissement gère l'organisation locale : classes, "
            "années académiques, comptes professeurs et étudiants, imports de listes.",
            {},
        ),
        (
            "Le professeur crée et suit les examens : énoncés, tests unitaires, "
            "publication, consultation des copies, lecture du journal d'incidents et "
            "ajustement éventuel des notes.",
            {},
        ),
        (
            "L'étudiant compose : connexion avec identifiant, mot de passe et code "
            "d'examen, acceptation du plein écran, salle d'attente, écriture, exécution "
            "d'essai, soumission.",
            {},
        ),
        (
            "Le surveillant physique n'a pas de compte applicatif. Il intervient en salle "
            "pour ce que le système ne voit pas.",
            {},
        ),
        (
            "Le système lui-même peut être vu comme acteur technique lorsqu'il agit de "
            "façon autonome, par exemple lorsqu'une soumission est forcée à l'expiration "
            "du temps, lorsqu'un incident déclenche une exclusion, ou lorsqu'une "
            "correction automatique est lancée après dépôt de copie.",
            {},
        ),
    ]

    ch23 = [
        (
            "Les besoins fonctionnels sont organisés par acteur.",
            {},
        ),
        (
            "Côté Super Admin : gérer les établissements, suivre l'activité plateforme et "
            "administrer les plans associés.",
            {},
        ),
        (
            "Côté administrateur : créer l'espace établissement, définir les classes, "
            "créer une année académique, importer les étudiants, gérer les comptes "
            "professeurs et étudiants (modification, désactivation, réactivation, "
            "réinitialisation de mot de passe).",
            {},
        ),
        (
            "Côté professeur : créer un examen, ajouter des exercices de code et des QCM, "
            "définir des tests unitaires, publier l'examen, suivre les participations, "
            "consulter les copies, lire le journal d'incidents, ajuster les notes.",
            {},
        ),
        (
            "Côté étudiant : se connecter avec identifiant, mot de passe et code "
            "d'examen, accepter le plein écran, patienter en salle d'attente, écrire et "
            "exécuter son code, soumettre avant l'expiration du temps.",
            {},
        ),
        (
            "Côté système : démarrer selon la logique temporelle retenue, sauvegarder "
            "régulièrement le code, traiter les incidents selon les règles prévues, et "
            "corriger automatiquement après soumission lorsque le mode le permet.",
            {},
        ),
    ]

    ch24 = [
        (
            "Au-delà des fonctions, plusieurs exigences transversales conditionnent la "
            "fiabilité du système. La sécurité doit rendre visible et coûteux le "
            "contournement de l'environnement de composition. La fiabilité doit limiter "
            "la perte de copie en cas de fermeture brutale ou de coupure courte. La "
            "performance doit supporter un pic de soumissions. La facilité d'usage le "
            "jour J doit rester une priorité. La compatibilité vise Chrome et Edge sur "
            "poste de travail. La résilience réseau doit permettre une reprise sans "
            "pénalité injustifiée. L'isolation de la sandbox doit empêcher l'accès réseau "
            "depuis le code exécuté. L'isolation multi établissement doit empêcher tout "
            "accès croisé aux données.",
            {},
        ),
        (
            "Ces exigences sont formulées comme des critères vérifiables, susceptibles "
            "d'être testés lors de la recette.",
            {},
        ),
    ]

    ch25 = [
        (
            "Plusieurs règles métier structurent le fonctionnement. Le code d'accès "
            "d'un examen est généré dans un format lisible, sans caractères ambigus. Un "
            "examen publié devient la référence de passation ; une fois l'épreuve "
            "engagée, le contenu n'est plus librement modifiable. Les sorties répétées "
            "du plein écran entraînent un avertissement puis une exclusion technique "
            "avec soumission forcée. Un étudiant ne voit pas son score immédiatement "
            "après dépôt. Les identifiants sont structurés selon le rôle, ce qui "
            "permet de déduire ce rôle à la connexion. Les données de chaque "
            "établissement restent isolées. La plateforme prévoit aussi une limite "
            "forte sur le nombre d'administrateurs actifs par établissement.",
            {},
        ),
    ]

    ch26 = [
        (
            "Plusieurs scénarios concrets illustrent l'usage attendu. Le professeur "
            "prépare l'examen à l'avance, définit exercices et tests, puis publie "
            "l'épreuve. Le jour J, il communique le code d'accès. Les étudiants se "
            "connectent et rejoignent la salle d'attente. Au démarrage, ils passent en "
            "composition. Pendant l'épreuve, un changement d'onglet est journalisé ; "
            "une sortie du plein écran déclenche d'abord un avertissement, puis une "
            "exclusion au seuil prévu. Une coupure réseau courte ne doit pas effacer "
            "le travail déjà sauvegardé. À l'expiration du temps, la copie en cours "
            "est soumise automatiquement. Après l'épreuve, le professeur lit le "
            "journal avant de communiquer les résultats.",
            {},
        ),
    ]

    ch27 = [
        (
            "Le cahier des charges de Cylentic est posé : périmètre du MVP, acteurs, "
            "besoins fonctionnels et non fonctionnels, règles métier et scénarios "
            "prioritaires. Sur cette base, la conception du chapitre suivant peut se "
            "développer sans improvisation fonctionnelle.",
            {},
        ),
    ]

    sections = [
        ("Résumé", resume_paras, ["Abstract"]),
        ("Abstract", abstract_paras, ["Introduction générale"]),
        ("Introduction générale", intro_paras, ["Chapitre 1"]),
        ("1.1 Contexte pédagogique et technologique", ch11, ["1.2 "]),
        ("1.2 Problématique", ch12, ["1.3 "]),
        ("1.3 Objectifs et hypothèse de travail", ch13, ["1.4 "]),
        (
            "1.4 État de l'art et comparaison des solutions existantes",
            ch14,
            ["1.5 "],
        ),
        ("1.5 Positionnement de Cylentic", ch15, ["1.6 "]),
        ("1.6 Conclusion partielle", ch16, ["Chapitre 2"]),
        ("2.1 Périmètre du projet et du MVP", ch21, ["2.2 "]),
        ("2.2 Identification des acteurs", ch22, ["2.3 "]),
        ("2.3 Besoins fonctionnels", ch23, ["2.4 "]),
        ("2.4 Besoins non fonctionnels", ch24, ["2.5 "]),
        ("2.5 Règles métier structurantes", ch25, ["2.6 "]),
        ("2.6 Scénarios d'usage prioritaires", ch26, ["2.7 "]),
        ("2.7 Conclusion partielle", ch27, ["Chapitre 3"]),
    ]

    for heading, items, stops in sections:
        fill_section(doc, heading, items, stops)

    liste = find_para(doc, lambda t: t.startswith("Liste des tableaux"))
    if liste is not None:
        has = False
        started = False
        for p in doc.paragraphs:
            if p._p is liste._p:
                started = True
                continue
            if not started:
                continue
            if p.text.strip() in {"Résumé", "Abstract", "Introduction générale"} or p.text.strip().startswith(
                "Résumé"
            ):
                break
            if "Tableau 1.1" in p.text:
                has = True
                break
        if not has:
            insert_paragraph_after(
                liste,
                "Tableau 1.1 : Comparaison des solutions existantes et de Cylentic",
                align="left",
                first_line=False,
                space_after=3,
            )

    full = "\n".join(p.text for p in doc.paragraphs)
    if "—" in full:
        raise SystemExit("Em dash detected")

    left = [
        p.text
        for p in doc.paragraphs
        if "à coller ici" in p.text
        and (
            p.text.startswith("[Contenu du point 1.")
            or p.text.startswith("[Contenu du point 2.")
        )
    ]
    doc.save(FINAL)
    print("Saved", FINAL)
    print("remaining ch1/2 placeholders:", left)
    print("tables:", len(doc.tables))
    print("has cinq chapitres:", "cinq chapitres" in full)
    print("Judge0 mentions:", full.count("Judge0"))


if __name__ == "__main__":
    main()
